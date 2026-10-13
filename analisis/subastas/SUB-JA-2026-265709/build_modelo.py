# -*- coding: utf-8 -*-
"""
Genera el modelo financiero en Excel de la subasta SUB-JA-2026-265709
(plaza de aparcamiento para moto nº 10, C/ Virtud 38 - C/ Constancia 71, Sevilla).

Uso:  python build_modelo.py [ruta_salida.xlsx]

Todas las cifras calculadas son fórmulas de Excel; los datos de partida
(documentos del expediente, comparables y supuestos) están en celdas de
entrada documentadas.
"""
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else "Modelo_Financiero_SUB-JA-2026-265709.xlsx"

# --------------------------------------------------------------------------
# Estilos
# --------------------------------------------------------------------------
FONT = "Arial"
BLUE, BLACK, GREEN, WHITE, GREY = "0000FF", "000000", "008000", "FFFFFF", "595959"
NAVY = "1F3864"
F_TITLE = Font(name=FONT, size=16, bold=True, color=NAVY)
F_SUB = Font(name=FONT, size=10, italic=True, color=GREY)
F_HDR = Font(name=FONT, size=10, bold=True, color=WHITE)
F_SEC = Font(name=FONT, size=11, bold=True, color=NAVY)
F_LBL = Font(name=FONT, size=10, color=BLACK)
F_LBL_B = Font(name=FONT, size=10, bold=True, color=BLACK)
F_IN = Font(name=FONT, size=10, color=BLUE)
F_FX = Font(name=FONT, size=10, color=BLACK)
F_FX_B = Font(name=FONT, size=10, bold=True, color=BLACK)
F_LINK = Font(name=FONT, size=10, color=GREEN)
F_LINK_B = Font(name=FONT, size=10, bold=True, color=GREEN)
F_NOTE = Font(name=FONT, size=9, italic=True, color=GREY)
F_KPI = Font(name=FONT, size=12, bold=True, color=GREEN)

FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_SEC = PatternFill("solid", fgColor="D9E1F2")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_KPI = PatternFill("solid", fgColor="F2F2F2")
FILL_OK = PatternFill("solid", fgColor="C6EFCE")
FILL_BAD = PatternFill("solid", fgColor="FFC7CE")
FILL_WARN = PatternFill("solid", fgColor="FFEB9C")

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOPLINE = Border(top=Side(style="thin", color=BLACK))

EUR0 = '#,##0" €";(#,##0" €");"-"'
EUR2 = '#,##0.00" €";(#,##0.00" €");"-"'
PCT1 = '0.0%;(0.0%);"-"'
PCT3 = '0.000%'
NUM0 = '#,##0;(#,##0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
IDX = '0'
MULT = '0.00"x"'
DATE = 'dd/mm/yyyy'
MONTH = 'mmm-yy'

WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()


def add_name(name, ws, cell):
    col = "".join(c for c in cell if c.isalpha())
    row = "".join(c for c in cell if c.isdigit())
    wb.defined_names[name] = DefinedName(name, attr_text=f"'{ws.title}'!${col}${row}")


def put(ws, cell, value, font=F_FX, fmt=None, fill=None, align=None, border=None, comment=None):
    c = ws[cell]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = border
    if comment:
        c.comment = Comment(comment, "Modelo")
    return c


def title(ws, text, sub=None):
    put(ws, "B2", text, F_TITLE)
    if sub:
        put(ws, "B3", sub, F_SUB)


def section(ws, row, text, first="B", last="H"):
    for col in range(ord(first), ord(last) + 1):
        ws[f"{chr(col)}{row}"].fill = FILL_SEC
    put(ws, f"{first}{row}", text, F_SEC, fill=FILL_SEC)


def header(ws, row, cols_values):
    for col, v in cols_values:
        put(ws, f"{col}{row}", v, F_HDR, fill=FILL_HDR, align=CENTER)


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def setup_print(ws, landscape=True):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# Fórmulas cerradas reutilizables (impuesto de adquisición y puja máxima)
def tax_expr(bid, reg="Regimen"):
    return (f"IF({reg}=1,TipoITP*MAX({bid},ValorReferencia),"
            f"TipoIVA*{bid}*(1-IVADeducible))")


def bid_from_avail(avail, reg="Regimen"):
    a = f"({avail})"
    return (f"MAX(0,IF({reg}=1,IF({a}/(1+TipoITP)>=ValorReferencia,{a}/(1+TipoITP),"
            f"{a}-TipoITP*ValorReferencia),{a}/(1+TipoIVA*(1-IVADeducible))))")


SALE_NET = "PrecioVenta*(1-DescuentoNeg)*(1-ComisionVenta*(1+TipoIVA))-Plusvalia"


def sale_net_of(price):
    return f"{price}*(1-DescuentoNeg)*(1-ComisionVenta*(1+TipoIVA))-Plusvalia"


def noi_of(rent, cuota="CuotaComunidad"):
    return (f"{rent}/(1+TipoIVA*RentaInclIVA)*12*(1-Desocupacion)*(1-GestionPct-MantenimientoPct)"
            f"-{cuota}*12-IBIAnual-SeguroAnual")


# ==========================================================================
# Hojas (orden final del libro)
# ==========================================================================
ws_res = wb.active
ws_res.title = "Resumen"
ws_sub = wb.create_sheet("Subasta")
ws_sup = wb.create_sheet("Supuestos")
ws_cmp = wb.create_sheet("Comparables")
ws_flip = wb.create_sheet("Compraventa")
ws_alq = wb.create_sheet("Alquiler")
ws_pm = wb.create_sheet("Puja máxima")
ws_sen = wb.create_sheet("Sensibilidad")
ws_esc = wb.create_sheet("Escenarios")
ws_rsk = wb.create_sheet("Riesgos")
ws_src = wb.create_sheet("Fuentes")

# ==========================================================================
# SUBASTA (datos de los documentos)
# ==========================================================================
ws = ws_sub
title(ws, "Datos de la subasta y del bien",
      "Extraídos de los documentos del expediente (edicto, certificación registral, decreto y diligencia de valoración).")
header(ws, 5, [("B", "Campo"), ("C", "Dato"), ("D", "Fuente"), ("E", "Estado")])

sub_rows = [
    ("SEC", "Identificación de la subasta"),
    ("Identificador", "SUB-JA-2026-265709", "Portal de Subastas del BOE", "Verificado", None),
    ("Tipo de subasta", "Judicial — vía de apremio (ejecución de títulos judiciales)", "Doc. 1 — Edicto", "Verificado", None),
    ("Órgano gestor", "Juzgado de Primera Instancia nº 4 de Sevilla · C/ Vermondo Resta s/n, 41018", "Doc. 1", "Verificado", None),
    ("Procedimiento", "Ejecución de títulos judiciales 1455/2022 · Negociado 2E · NIG 4109142120190033606", "Doc. 1", "Verificado", None),
    ("Ejecutante", "Comunidad de Propietarios Calle Virtud 44", "Doc. 1", "Verificado", None),
    ("Ejecutado (titular registral)", "SUPERCHINDO, S.L. (CIF B92548411)", "Docs. 1 y 2", "Verificado", None),
    ("Fecha del edicto", datetime(2025, 6, 16), "Doc. 1", "Verificado", DATE),
    ("Fecha de cierre de la subasta", "=FechaCierre", "No figura en los PDF · estimada", "PENDIENTE", DATE),
    ("Lotes", "Lote único", "Doc. 1", "Verificado", None),
    ("SEC", "Cantidades reclamadas"),
    ("Principal reclamado (edicto)", 5130.01, "Doc. 1 — Edicto (16/06/2025)", "Verificado", EUR2),
    ("Intereses y costas presupuestados", 4953.72, "Doc. 1", "Verificado", EUR2),
    ("Total reclamado", "=C{principal}+C{intereses}", "Cálculo", "", EUR2),
    ("Embargo anotado (letra A, 30/12/2022)", "6.476,23 € de principal + 4.150,93 € de intereses y costas", "Doc. 2 — Certificación registral", "Verificado", None),
    ("Mandamiento de certificación (15/03/2023)", "13.836,42 € de principal + 4.150,93 € de intereses y costas", "Doc. 2", "Verificado", None),
    ("SEC", "Descripción del bien"),
    ("Tipo de bien", "Plaza de aparcamiento nº 10 en planta sótano (apta para moto)", "Docs. 1, 2, 4 y 5", "Verificado", None),
    ("Dirección", "C/ Virtud 38 y C/ Constancia 71 (acceso por C/ Constancia), Sevilla", "Doc. 2", "Verificado", None),
    ("Barrio / código postal", "López de Gomara (Triana) · 41010", "Idealista / RealAdvisor", "Fuente externa", None),
    ("Superficie construida (m²)", 9.50, "Doc. 2 — Certificación registral", "Verificado", NUM2),
    ("Superficie útil (m²)", 3.75, "Doc. 2 (el edicto solo cita la construida)", "Verificado", NUM2),
    ("Linderos", "Frente: zona de maniobra · Izquierda: trastero nº 17 · Fondo: casa nº 69 de C/ Constancia · Derecha: zona de maniobra", "Doc. 2", "Verificado", None),
    ("Cuota de participación en la comunidad", 0.00562, "Doc. 2", "Verificado", PCT3),
    ("Finca registral", "Nº 37.841 de Sevilla · Registro de la Propiedad nº 2 · Tomo 1.844, Libro 1.175, Folio 181", "Doc. 2", "Verificado", None),
    ("IDUFIR", "41028000513716", "Doc. 2", "Verificado", None),
    ("Título del ejecutado", "División horizontal, escritura de 12/06/2018 (protocolo 1503); inscripción 1ª de 16/11/2018", "Doc. 2", "Verificado", None),
    ("Referencia catastral", "No consta en los documentos", "Sede Electrónica del Catastro", "PENDIENTE", None),
    ("SEC", "Valoración y condiciones económicas"),
    ("Valoración pericial definitiva", 7078.45, "Doc. 4 — Decreto 19/03/2025 (sin alegaciones)", "Verificado", EUR2),
    ("Cargas anteriores deducidas", 0, "Doc. 5 — Diligencia 19/03/2025", "Verificado", EUR2),
    ("Valor de subasta", 7078.45, "Doc. 5 — Diligencia 19/03/2025 · Doc. 1", "Verificado", EUR2),
    ("Valor de subasta por m² construido", "=C{vs}/C{supc}", "Cálculo", "", EUR0),
    ("Depósito para pujar (5%)", "=Deposito", "Art. 669.1 LEC", "Calculado", EUR2),
    ("70% del valor de subasta", "=UmbralSetenta", "Art. 670 LEC", "Calculado", EUR2),
    ("50% del valor de subasta", "=UmbralCincuenta", "Art. 670.4 LEC", "Calculado", EUR2),
    ("Puja mínima", "Sin puja mínima (rige el art. 670 LEC)", "Doc. 1, condición particular 3ª", "Verificado", None),
    ("Tramos entre pujas", "Consultar en la ficha del Portal", "Portal de Subastas", "PENDIENTE", None),
    ("Pago del resto del precio", "40 días desde la aprobación del remate", "Art. 670.1 LEC", "Verificado", None),
    ("Reserva de postura / cesión de remate", "Reserva permitida (art. 652 LEC); solo el ejecutante puede ceder el remate (art. 647.3 LEC)", "Doc. 1", "Verificado", None),
    ("SEC", "Situación jurídica"),
    ("Cargas registrales", "(1) Notas de afección fiscal · (2) Embargo letra A a favor de la Comunidad ejecutante (se cancela con la adjudicación) · (3) Prórroga letra B por 4 años desde 24/03/2023", "Doc. 2", "Verificado", None),
    ("Titulares de derechos posteriores", "Ninguno (no proceden comunicaciones del art. 659 LEC)", "Doc. 2", "Verificado", None),
    ("Fecha de la certificación de cargas", datetime(2023, 3, 24), "Doc. 2", "Actualizar", DATE),
    ("Situación posesoria", "No consta (condición general 4ª del edicto)", "Doc. 1", "PENDIENTE", None),
    ("Visitas", "Se puede pedir al Juzgado la inspección del bien durante la licitación (art. 669.3 LEC)", "Doc. 1", "Verificado", None),
    ("Otros bienes del procedimiento", "El edicto menciona un local comercial \"vacío y sin actividad\" (posible subasta relacionada)", "Doc. 1", "Revisar", None),
    ("SEC", "Documentos analizados"),
    ("Documento 1", "Edicto de subasta del JPI nº 4 de Sevilla, 16/06/2025 (6 págs.)", "Aportado", "Analizado", None),
    ("Documento 2", "Mandamiento (15/03/2023) y certificación de dominio y cargas del Registro nº 2 de Sevilla (24/03/2023) (10 págs.)", "Aportado", "Analizado", None),
    ("Documento 3", "No aportado", "—", "Falta", None),
    ("Documento 4", "Decreto de valoración definitiva, 19/03/2025 (3 págs.)", "Aportado", "Analizado", None),
    ("Documento 5", "Diligencia de ordenación: valoración a efectos de subasta, 19/03/2025 (2 págs.)", "Aportado", "Analizado", None),
]
SUB_KEYS = {"Principal reclamado (edicto)": "principal", "Intereses y costas presupuestados": "intereses",
            "Total reclamado": "total", "Superficie construida (m²)": "supc", "Superficie útil (m²)": "supu",
            "Valoración pericial definitiva": "tasacion", "Cargas anteriores deducidas": "cargas",
            "Valor de subasta": "vs"}
SUBROW = {}
r = 6
for item in sub_rows:
    if item[0] == "SEC":
        section(ws, r, item[1], "B", "E")
        r += 1
        continue
    lbl, val, src, status, fmt = item
    if lbl in SUB_KEYS:
        SUBROW[SUB_KEYS[lbl]] = r
    put(ws, f"B{r}", lbl, F_LBL)
    if isinstance(val, (int, float, datetime)):
        font = F_IN                      # dato documental (entrada)
    elif isinstance(val, str) and val.startswith("=C"):
        font = F_FX                      # cálculo dentro de la hoja
    elif isinstance(val, str) and val.startswith("="):
        font = F_LINK                    # vínculo a Supuestos
    else:
        font = F_LBL                     # texto
    put(ws, f"C{r}", val, font, fmt=fmt, align=Alignment(wrap_text=True, vertical="top", horizontal="left"), border=BOX,
        fill=FILL_KEY if status == "PENDIENTE" else None)
    put(ws, f"D{r}", src, F_NOTE, align=WRAP)
    st_font = Font(name=FONT, size=9, bold=True, color="C00000" if status in ("PENDIENTE", "Falta") else GREEN if status in ("Verificado", "Analizado") else GREY)
    put(ws, f"E{r}", status, st_font)
    r += 1
# Resolver marcadores {clave} en fórmulas de la propia hoja
for row in ws.iter_rows(min_row=6, max_row=r, min_col=3, max_col=3):
    c = row[0]
    if isinstance(c.value, str) and c.value.startswith("=") and "{" in c.value:
        c.value = c.value.format(**SUBROW)
add_name("SupConstruida", ws, f"C{SUBROW['supc']}")
add_name("SupUtil", ws, f"C{SUBROW['supu']}")
add_name("CargasAnteriores", ws, f"C{SUBROW['cargas']}")
add_name("Tasacion", ws, f"C{SUBROW['tasacion']}")
widths(ws, {"A": 2, "B": 40, "C": 78, "D": 40, "E": 14})
setup_print(ws)

# ==========================================================================
# SUPUESTOS
# ==========================================================================
ws = ws_sup
title(ws, "Supuestos del modelo — Subasta SUB-JA-2026-265709",
      "Plaza de aparcamiento (moto) nº 10 · C/ Virtud 38 – C/ Constancia 71 · López de Gomara (Triana), Sevilla")
put(ws, "B4", "Leyenda:", F_LBL_B)
put(ws, "C4", "Entrada", F_IN, border=BOX)
put(ws, "D4", "Fórmula", F_FX, border=BOX)
put(ws, "E4", "Vínculo", F_LINK, border=BOX)
put(ws, "F4", "Clave / verificar", F_IN, fill=FILL_KEY, border=BOX)
put(ws, "H4", "Solo se editan las celdas azules. La columna F contiene siempre el valor que usa el modelo.", F_NOTE)

r = 6
section(ws, r, "1. Selección de escenario y estrategia")
rows = [
    ("Escenario activo (1 = Pesimista · 2 = Base · 3 = Optimista)", 2, "Escenario", "nº", "Cambia todas las variables de la tabla 2 a la vez.", True),
    ("Nombre del escenario", '=CHOOSE(Escenario,"Pesimista","Base","Optimista")', "NombreEscenario", "", "", False),
    ("Estrategia principal (1 = Compraventa · 2 = Alquiler)", 1, "Estrategia", "nº", "Determina la puja máxima recomendada en el Resumen.", True),
    ("Nombre de la estrategia", '=CHOOSE(Estrategia,"Compraventa (reventa)","Alquiler (patrimonial)")', "NombreEstrategia", "", "", False),
]
r += 1
for lbl, val, nm, unit, note, is_in in rows:
    put(ws, f"B{r}", lbl, F_LBL)
    put(ws, f"F{r}", val, F_IN if is_in else F_FX, fmt=IDX if is_in else None,
        fill=FILL_KEY if is_in else None, border=BOX)
    put(ws, f"G{r}", unit, F_LBL)
    put(ws, f"H{r}", note, F_NOTE)
    add_name(nm, ws, f"F{r}")
    r += 1
dv_esc = DataValidation(type="list", formula1='"1,2,3"', allow_blank=False)
dv_est = DataValidation(type="list", formula1='"1,2"', allow_blank=False)
dv_01 = DataValidation(type="list", formula1='"0,1"', allow_blank=False)
ws.add_data_validation(dv_esc)
ws.add_data_validation(dv_est)
ws.add_data_validation(dv_01)
dv_esc.add("F7")
dv_est.add("F9")

r += 1
section(ws, r, "2. Variables por escenario")
r += 1
header(ws, r, [("B", "Variable"), ("C", "Pesimista"), ("D", "Base"), ("E", "Optimista"),
               ("F", "Activo"), ("G", "Unidad"), ("H", "Justificación / fuente")])
r += 1
SCN = {}  # nombre -> fila de la tabla de escenarios
scen_rows = [
    ("Puja (precio de remate)", 6500, 5500, 5000, "Puja", EUR0, "€",
     "Base 77,7% del valor de subasta (> umbral 70% del art. 670 LEC). Ver hoja 'Puja máxima'."),
    ("Precio de salida a la venta", 8000, 10000, 12000, "PrecioVenta", EUR0, "€",
     "Mediana oferta plazas de moto Triana/Remedios 11.900 €; mismo edificio 10.000 € (hoja Comparables)."),
    ("Renta mensual de mercado (lo que paga el inquilino)", 35, 45, 50, "Renta", EUR0, "€/mes",
     "Comparables Triana: 40–55 €/mes; C/ Ardilla 28 (mismo barrio) 40 €/mes."),
    ("Mes de toma de posesión (desde el cierre)", 6, 4, 3, "MesPosesion", IDX, "mes",
     "Situación posesoria 'no consta'. Pesimista: ocupada y posesión por art. 675 LEC."),
    ("Coste de toma de posesión", 800, 100, 60, "CostePosesion", EUR0, "€",
     "Mando/llave del garaje; pesimista incluye abogado y procurador para el lanzamiento."),
    ("Cuota de comunidad", 15, 10, 8, "CuotaComunidad", EUR0, "€/mes",
     "Estimación con cuota de participación 0,562%. Pedir certificado al administrador."),
    ("Meses de cuotas impagadas a cargo del adjudicatario", 46, 36, 12, "MesesDeuda", IDX, "meses",
     "Afección real: año en curso + 3 anteriores (art. 9.1.e LPH). Máx. ≈ 46 meses a oct-2026."),
    ("Meses de comercialización hasta la venta", 8, 4, 2, "MesesVenta", IDX, "meses",
     "Mercado de plazas de moto estrecho; plazos de venta largos."),
    ("Régimen fiscal de la adquisición (1 = ITP · 2 = IVA)", 2, 1, 1, "Regimen", IDX, "nº",
     "El titular es una S.L. por división horizontal (2018): riesgo de 'primera entrega' con IVA 21%."),
]
for lbl, p, b, o, nm, fmt, unit, note in scen_rows:
    put(ws, f"B{r}", lbl, F_LBL)
    for col, v in (("C", p), ("D", b), ("E", o)):
        put(ws, f"{col}{r}", v, F_IN, fmt=fmt, border=BOX, fill=FILL_KEY if col == "D" else None)
    put(ws, f"F{r}", f"=INDEX(C{r}:E{r},Escenario)", F_FX_B, fmt=fmt, border=BOX)
    put(ws, f"G{r}", unit, F_LBL)
    put(ws, f"H{r}", note, F_NOTE)
    add_name(nm, ws, f"F{r}")
    SCN[nm] = r
    if nm == "Regimen":
        dv_est.add(f"C{r}:E{r}")
    r += 1


def block(rows_def, start):
    """rows_def: (label, value, name, fmt, unit, note, kind) kind in in/key/fx/link"""
    rr = start
    for lbl, val, nm, fmt, unit, note, kind in rows_def:
        put(ws, f"B{rr}", lbl, F_LBL_B if kind == "fxb" else F_LBL)
        font = {"in": F_IN, "key": F_IN, "fx": F_FX, "fxb": F_FX_B, "link": F_LINK}[kind]
        put(ws, f"F{rr}", val, font, fmt=fmt, border=BOX,
            fill=FILL_KEY if kind == "key" else None)
        put(ws, f"G{rr}", unit, F_LBL)
        put(ws, f"H{rr}", note, F_NOTE)
        if nm:
            add_name(nm, ws, f"F{rr}")
        rr += 1
    return rr


r += 1
section(ws, r, "3. Datos de la subasta (vinculados a la hoja Subasta)")
r = block([
    ("Valor de subasta", f"=Subasta!C{SUBROW['vs']}", "ValorSubasta", EUR2, "€", "Diligencia de 19/03/2025 (cargas deducibles 0 €).", "link"),
    ("Depósito para pujar (% del valor de subasta)", 0.05, "DepositoPct", PCT1, "%", "Art. 669.1 LEC.", "in"),
    ("Depósito para pujar", "=ValorSubasta*DepositoPct", "Deposito", EUR2, "€", "Se consigna a través del Portal de Subastas (AEAT).", "fx"),
    ("Umbral de aprobación directa (% del valor)", 0.70, "PctAprobacion", PCT1, "%", "Art. 670.1 LEC.", "in"),
    ("Umbral de aprobación directa", "=ValorSubasta*PctAprobacion", "UmbralSetenta", EUR2, "€", "Puja ≥ este importe: se aprueba el remate al mejor postor.", "fx"),
    ("Umbral mínimo de referencia (% del valor)", 0.50, "PctMinimo", PCT1, "%", "Art. 670.4 LEC.", "in"),
    ("Umbral mínimo de referencia", "=ValorSubasta*PctMinimo", "UmbralCincuenta", EUR2, "€", "Por debajo, el remate queda a criterio del LAJ.", "fx"),
    ("Cantidad reclamada (principal + intereses y costas)", f"=Subasta!C{SUBROW['total']}", "CantidadReclamada", EUR2, "€", "Edicto de 16/06/2025.", "link"),
], r + 1)

r += 1
section(ws, r, "4. Fiscalidad de la adquisición")
r = block([
    ("Régimen aplicado", '=IF(Regimen=1,"ITP — segunda transmisión","IVA — primera entrega del promotor")', None, None, "", "Se elige en la tabla 2 (por escenario).", "fx"),
    ("Tipo ITP Andalucía (tipo general)", 0.07, "TipoITP", PCT1, "%", "Ley 5/2021 de Tributos Cedidos de Andalucía.", "in"),
    ("Valor de referencia catastral (estimado)", 7078.45, "ValorReferencia", EUR2, "€",
     "PENDIENTE: consultar en Sede del Catastro. Proxy = valoración pericial. Base ITP = máx(remate; valor de referencia).", "key"),
    ("Tipo IVA (plaza de garaje no anexa a vivienda)", 0.21, "TipoIVA", PCT1, "%", "Art. 90 LIVA. También grava la comisión de venta y el alquiler de la plaza.", "in"),
    ("¿IVA soportado deducible? (1 = Sí · 0 = No)", 0, "IVADeducible", IDX, "nº", "Sí solo si el comprador es empresario y alquila la plaza con IVA.", "in"),
    ("Base imponible", "=IF(Regimen=1,MAX(Puja,ValorReferencia),Puja)", "BaseImponible", EUR0, "€", "", "fx"),
    ("Impuesto de adquisición", "=" + tax_expr("Puja"), "ImpuestoAdq", EUR0, "€", "", "fxb"),
], r + 1)
dv_01.add(f"F{r-3}")

r += 1
section(ws, r, "5. Otros costes de adquisición")
r = block([
    ("Registro de la Propiedad (inscripción + cancelación de anotaciones A y B)", 150, "Registro", EUR0, "€", "Arancel registral para un valor ≈ 7.000 €.", "in"),
    ("Gestoría / liquidación de impuestos", 250, "Gestoria", EUR0, "€", "IVA incluido.", "in"),
    ("Abogado y procurador (opcional)", 0, "Abogado", EUR0, "€", "No son obligatorios para pujar como tercero.", "in"),
    ("Deuda de comunidad a cargo del adjudicatario", "=CuotaComunidad*MesesDeuda", "DeudaComunidad", EUR0, "€", "Cuota × meses impagados afectos (tabla 2).", "fx"),
    ("IBI anual estimado", 40, "IBIAnual", EUR0, "€/año", "PENDIENTE: consultar recibo en la Agencia Tributaria de Sevilla.", "key"),
    ("Años de IBI pendientes a cargo del adjudicatario", 4, "AnosIBI", IDX, "años", "Afección del inmueble (art. 64 TRLRHL), máximo 4 años.", "in"),
    ("Deuda de IBI a cargo del adjudicatario", "=IBIAnual*AnosIBI", "DeudaIBI", EUR0, "€", "", "fx"),
    ("Coste de toma de posesión", "=CostePosesion", None, EUR0, "€", "Tabla 2.", "fx"),
    ("Total de costes fijos de adquisición", "=Registro+Gestoria+Abogado+DeudaComunidad+DeudaIBI+CostePosesion", "CostesFijosAdq", EUR0, "€", "Excluye la puja y el impuesto.", "fxb"),
], r + 1)

r += 1
section(ws, r, "6. Calendario (meses desde el cierre de la subasta)")
r = block([
    ("Fecha de cierre de la subasta", datetime(2026, 10, 20), "FechaCierre", DATE, "fecha",
     "PENDIENTE: no figura en los PDF. Estimada por la secuencia de identificadores del Portal de Subastas.", "key"),
    ("Mes de pago del resto del precio", 1, "MesPago", IDX, "mes", "Hasta 40 días desde la aprobación del remate (art. 670.1 LEC).", "in"),
    ("Mes del decreto de adjudicación y testimonio", 3, "MesDecreto", IDX, "mes", "Plazo habitual en los juzgados de Sevilla: 2–6 meses.", "in"),
    ("Mes de toma de posesión", "=MesPosesion", None, IDX, "mes", "Tabla 2.", "fx"),
    ("Mes de venta (estrategia compraventa)", "=MesPosesion+MesesVenta", "MesVenta", IDX, "mes", "Máximo modelizado: mes 23.", "fx"),
    ("Meses hasta alquilar tras la posesión", 1, "MesesPuestaAlquiler", IDX, "meses", "", "in"),
    ("Mes de inicio del alquiler", "=MesPosesion+MesesPuestaAlquiler", "MesAlquiler", IDX, "mes", "", "fx"),
], r + 1)

r += 1
section(ws, r, "7. Gastos corrientes de tenencia")
r = block([
    ("Cuota de comunidad", "=CuotaComunidad", None, EUR0, "€/mes", "Tabla 2.", "fx"),
    ("Seguro anual", 0, "SeguroAnual", EUR0, "€/año", "Cubierto por la póliza de la comunidad.", "in"),
    ("Gasto mensual de tenencia", "=CuotaComunidad+(IBIAnual+SeguroAnual)/12", "GastoMensual", EUR2, "€/mes", "", "fx"),
    ("Meses de tenencia (compraventa)", "=MAX(0,MesVenta-MesDecreto)", "MesesTenenciaFlip", IDX, "meses", "Desde el decreto hasta la venta.", "fx"),
    ("Coste de tenencia (compraventa)", "=GastoMensual*MesesTenenciaFlip", "CosteTenenciaFlip", EUR0, "€", "", "fx"),
    ("Meses de tenencia hasta alquilar", "=MAX(0,MesAlquiler-MesDecreto)", "MesesTenenciaAlq", IDX, "meses", "", "fx"),
    ("Coste de tenencia hasta alquilar", "=GastoMensual*MesesTenenciaAlq", "CosteTenenciaAlq", EUR0, "€", "", "fx"),
], r + 1)

r += 1
section(ws, r, "8. Venta (estrategia compraventa)")
r = block([
    ("Precio de salida", "=PrecioVenta", None, EUR0, "€", "Tabla 2.", "fx"),
    ("Descuento de negociación", 0.05, "DescuentoNeg", PCT1, "%", "Los comparables son precios de oferta, no de cierre.", "in"),
    ("Precio de cierre", "=PrecioVenta*(1-DescuentoNeg)", "PrecioCierre", EUR0, "€", "", "fx"),
    ("Comisión de intermediación (sobre precio, + IVA)", 0.03, "ComisionVenta", PCT1, "%", "Poner 0% si se vende como particular.", "in"),
    ("Plusvalía municipal (IIVTNU) en la venta", 70, "Plusvalia", EUR0, "€", "Valor catastral del suelo muy bajo; tenencia < 1 año.", "in"),
    ("Ingreso neto por la venta", "=" + SALE_NET, "VentaNeta", EUR0, "€", "", "fxb"),
    ("Tipo impositivo sobre la ganancia", 0.19, "TipoGanancia", PCT1, "%", "IRPF base del ahorro (19% hasta 6.000 €). 25% si se compra con una S.L.", "in"),
], r + 1)

r += 1
section(ws, r, "9. Alquiler (estrategia patrimonial)")
r = block([
    ("Renta mensual de mercado", "=Renta", None, EUR0, "€/mes", "Tabla 2.", "fx"),
    ("¿La renta de mercado incluye IVA? (1 = Sí · 0 = No)", 1, "RentaInclIVA", IDX, "nº",
     "El alquiler de una plaza de garaje independiente está sujeto y no exento de IVA (art. 20.Uno.23 LIVA).", "in"),
    ("Renta mensual neta de IVA", "=Renta/(1+TipoIVA*RentaInclIVA)", "RentaNeta", EUR2, "€/mes", "", "fx"),
    ("Desocupación e impagos (% de ingresos)", 0.08, "Desocupacion", PCT1, "%", "≈ 1 mes al año.", "in"),
    ("Gestión del alquiler (% de ingresos)", 0.0, "GestionPct", PCT1, "%", "Gestión propia.", "in"),
    ("Reserva de mantenimiento (% de ingresos)", 0.02, "MantenimientoPct", PCT1, "%", "", "in"),
    ("Crecimiento anual de la renta", 0.025, "CrecRenta", PCT1, "%", "", "in"),
    ("Inflación anual de gastos", 0.025, "InflGastos", PCT1, "%", "", "in"),
    ("Revalorización anual del activo", 0.02, "Revalorizacion", PCT1, "%", "Sobre el precio de salida actual.", "in"),
    ("Plusvalía municipal en la venta final (año 10)", 150, "PlusvaliaSalida", EUR0, "€", "", "in"),
    ("Tasa de descuento (VAN)", 0.06, "TasaDescuento", PCT1, "%", "Coste de oportunidad del capital.", "in"),
    ("NOI año 1 (ingresos netos − gastos)", "=" + noi_of("Renta"), "NOIInicial", EUR0, "€/año", "", "fxb"),
], r + 1)
dv_01.add(f"F{r-11}")

r += 1
section(ws, r, "10. Objetivos de rentabilidad")
r = block([
    ("ROI mínimo exigido (compraventa)", 0.20, "ROIObjetivo", PCT1, "%", "Sobre la inversión total.", "key"),
    ("Rentabilidad neta mínima exigida (alquiler)", 0.05, "RentabObjetivo", PCT1, "%", "NOI año 1 / inversión total.", "key"),
], r + 1)

widths(ws, {"A": 2, "B": 58, "C": 12, "D": 12, "E": 12, "F": 16, "G": 9, "H": 95})
ws.freeze_panes = "C6"
setup_print(ws)

# ==========================================================================
# COMPARABLES
# ==========================================================================
ws = ws_cmp
title(ws, "Comparables de mercado — plazas de moto en Sevilla",
      "Precios de oferta publicados en idealista, consultados el 13/10/2026. Superficie en blanco = no publicada.")
URL = ("https://www.idealista.com/inmueble/{}/?utm_medium=generativeAI&utm_campaign=appChatgpt"
       "&utm_creation=link&utm_source=idGpt&utm_project=leadGeneration&utm_link=inChat_detail")

cmp_tables = [
    ("A. Venta — plazas de moto en Triana y Los Remedios", "€", [
        (112744803, "C/ Virtud 44 (acceso por C/ Constancia)", "López de Gomara (Triana)", "Moto · 1 plaza", 4, 10000,
         "MISMO EDIFICIO y misma superficie registral (9,50 m² const.): puede ser esta misma finca. Particular; posesión 'a partir de noviembre'. No es un comparable independiente."),
        (104181518, "C/ Troya", "Triana", "Moto · 2 motos", None, 13500, "Se puede alquilar."),
        (102202705, "C/ Virgen de la Estrella 10", "Los Remedios", "Moto · 1 plaza", 3, 11900, "Rampa, puerta automática."),
    ]),
    ("B. Venta — plazas de moto en el resto de Sevilla capital", "€", [
        (112644449, "C/ Francos", "Centro (Alfalfa)", "Moto · 1,83 × 2,10 m", 3.84, 16000, "Planta -1, zona sin aparcamiento."),
        (112732191, "C/ Juan Antonio Cavestany 20", "Santa Justa", "Moto o Smart · 4,40 × 2,10 m", 10, 15000, ""),
        (111806571, "C/ Mendigorría 5", "Alameda – San Lorenzo", "Moto · 1 plaza", None, 18000, "Vigilancia 24 h."),
        (112521156, "C/ Conde Negro 15", "Puerta Carmona", "2 motos", None, 12250, "Alquilada a 70 €/mes."),
        (111834117, "Parqueflores VI Fase", "Los Carteros – San Diego", "2–3 motos · 4,50 × 2,20 m", 9.9, 16000, ""),
        (100622617, "C/ Marqués del Nervión 1", "Nervión", "Motos / Smart (se venden juntas)", 6, 15500, ""),
        (105263661, "C/ Almenas 55", "Bellavista", "4 motos · 14,5 m² útiles", 14.5, 13000, "Alquiladas; anuncia rentabilidad del 8%."),
        (109368921, "C/ Marathón", "Avenida de las Ciencias", "Moto · 12 m²", 12, 14500, "Venta directa de promotora (impuestos aparte)."),
        (107119263, "C/ Luis Montoto 63", "Santa Justa", "Moto · 1 plaza", None, 8200, "Opción de 3 plazas por 24.000 €."),
        (112181517, "C/ Monzón 1 D", "Heliópolis", "2–3 motos", None, 8000, "Concesión administrativa (70 años)."),
        (105870147, "C/ Fernández de Ribera", "Nervión", "2–3 motos", None, 7000, "Planta -3, acceso complicado."),
        (111785329, "C/ Hispano Aviación 2", "Tablada", "Moto · 1 plaza", None, 5000, "Alquilable a 30 €/mes."),
        (106975276, "C/ Hispano Aviación 2", "Tablada", "Moto · 1 plaza", None, 4500, ""),
        (29257022, "C/ Electricidad 5", "Santa Aurelia", "2 motos", None, 5000, ""),
    ]),
    ("C. Alquiler — plazas de moto en Triana", "€/mes", [
        (82917753, "C/ Ardilla 28", "López de Gomara (Triana)", "Moto", None, 40, "El comparable más cercano (mismo barrio)."),
        (89222824, "C/ Pagés del Corro 82", "Ronda de Triana – Patrocinio", "Moto", None, 40, "Cámaras, puerta automática."),
        (96520950, "C/ Doctores Lasso Simarro", "Betis – Pagés del Corro", "Moto", None, 40, "Acceso por elevador."),
        (111286433, "C/ Alfarería 124", "Betis – Pagés del Corro", "Moto (parking profesional)", None, 49, "24/7, taquilla y seguro incluidos."),
        (84436084, "C/ Castilla 83", "Betis – Pagés del Corro", "Moto", None, 50, ""),
        (87211204, "Ronda de Triana 14", "Ronda de Triana – Patrocinio", "Moto", None, 55, "Aparcamiento Crucero Baleares."),
        (86864183, "C/ José León Sanz 2", "Barrio León", "Moto · caben 3", None, 55, ""),
    ]),
    ("D. Alquiler — plazas de moto en Los Remedios y Tablada", "€/mes", [
        (38430792, "C/ Juan Sebastián Elcano 15", "Los Remedios", "Moto", None, 45, "Videovigilancia."),
        (92276326, "C/ Virgen de las Huertas 2", "Plaza de Cuba", "Moto · caben 3", None, 45, "Parking Resitur."),
        (31113331, "C/ Salado 4", "Plaza de Cuba", "Moto pequeña/mediana", None, 45, "Todo incluido."),
        (84484073, "C/ Virgen de la Estrella 10", "Los Remedios", "Moto · 2 × 2,40 m", 4.8, 50, ""),
        (85885052, "C/ Virgen de la Estrella 10", "Los Remedios", "Moto", None, 30, ""),
        (93212844, "C/ Virgen de la Estrella 6", "Los Remedios", "Moto", None, 30, "Comunidad incluida."),
        (94516135, "C/ Virgen de la Estrella 2", "Los Remedios", "Moto", None, 25, ""),
        (103494976, "C/ Virgen del Refugio", "Parque de los Príncipes", "Moto (compartida)", None, 40, ""),
        (107688502, "C/ Hispano Aviación 2", "Tablada", "Moto", None, 40, ""),
        (109160656, "C/ Hispano Aviación 2", "Tablada", "Moto", None, 30, "Incluye semana de feria."),
        (97438138, "C/ Hispano Aviación 2", "Tablada", "Moto", None, 35, ""),
        (111143270, "C/ Hispano Aviación 2", "Tablada", "Moto", None, 30, ""),
        (105349791, "C/ Hispano Aviación 2", "Tablada", "Moto", None, 60, ""),
    ]),
]
r = 5
STATS = {}
for tname, unit, items in cmp_tables:
    section(ws, r, tname, "B", "J")
    r += 1
    header(ws, r, [("B", "Ref. idealista"), ("C", "Ubicación"), ("D", "Barrio"), ("E", "Tipo / capacidad"),
                   ("F", "Sup. (m²)"), ("G", f"Precio ({unit})"), ("H", f"{unit}/m²"), ("I", "Observaciones"), ("J", "Enlace")])
    r += 1
    first = r
    for code, loc, barrio, tipo, sup, price, obs in items:
        put(ws, f"B{r}", str(code), F_LBL)
        put(ws, f"C{r}", loc, F_LBL)
        put(ws, f"D{r}", barrio, F_LBL)
        put(ws, f"E{r}", tipo, F_LBL)
        put(ws, f"F{r}", sup, F_IN, fmt=NUM2)
        put(ws, f"G{r}", price, F_IN, fmt=EUR0)
        put(ws, f"H{r}", f'=IF(F{r}>0,G{r}/F{r},"")', F_FX, fmt=EUR0)
        put(ws, f"I{r}", obs, F_NOTE, align=WRAP)
        link = ws[f"J{r}"]
        link.value = URL.format(code)
        link.hyperlink = URL.format(code)
        link.font = Font(name=FONT, size=8, color="0563C1", underline="single")
        if code == 112744803:
            for col in "BCDEFGHI":
                ws[f"{col}{r}"].fill = FILL_WARN
        r += 1
    last = r - 1
    stats = [("Nº de comparables", f"=COUNT(G{first}:G{last})", IDX),
             ("Media", f"=AVERAGE(G{first}:G{last})", EUR0),
             ("Mediana", f"=MEDIAN(G{first}:G{last})", EUR0),
             ("Mínimo", f"=MIN(G{first}:G{last})", EUR0),
             ("Máximo", f"=MAX(G{first}:G{last})", EUR0)]
    for lbl, fx, fmt in stats:
        put(ws, f"F{r}", lbl, F_LBL_B, align=Alignment(horizontal="right"))
        put(ws, f"G{r}", fx, F_FX_B, fmt=fmt, border=BOX)
        if lbl == "Mediana":
            STATS[tname[0]] = f"G{r}"
        r += 1
    r += 1

section(ws, r, "E. Conclusión de valor", "B", "J")
r += 1
concl = [
    ("Mediana de venta — Triana y Los Remedios", f"={STATS['A']}", F_FX, EUR0),
    ("Mediana de venta — resto de Sevilla", f"={STATS['B']}", F_FX, EUR0),
    ("Valoración pericial del Juzgado", "=Tasacion", F_LINK, EUR0),
    ("Precio de salida usado en el modelo (escenario activo)", "=PrecioVenta", F_LINK, EUR0),
    ("Mediana de alquiler — Triana", f"={STATS['C']}", F_FX, EUR0),
    ("Mediana de alquiler — Los Remedios y Tablada", f"={STATS['D']}", F_FX, EUR0),
    ("Renta usada en el modelo (escenario activo)", "=Renta", F_LINK, EUR0),
    ("Rentabilidad bruta implícita a precio de mercado (renta × 12 / precio de salida)", "=Renta*12/PrecioVenta", F_FX, PCT1),
]
for lbl, fx, font, fmt in concl:
    put(ws, f"B{r}", lbl, F_LBL)
    put(ws, f"G{r}", fx, font, fmt=fmt, border=BOX)
    r += 1
r += 1
put(ws, f"B{r}", ("Criterio: la plaza solo admite una moto (3,75 m² útiles), así que el precio base (10.000 €) se sitúa por debajo de la mediana "
                  "de Triana/Remedios (11.900 €), que incluye plazas para dos motos, y en línea con el anuncio del mismo edificio. "
                  "Referencia no comparable: plazas de coche de obra nueva en C/ Trabajo (López de Gomara) a 35.000–42.500 € + IVA."), F_NOTE, align=WRAP)
ws.merge_cells(f"B{r}:J{r}")
ws.row_dimensions[r].height = 42
widths(ws, {"A": 2, "B": 13, "C": 32, "D": 27, "E": 28, "F": 13, "G": 13, "H": 10, "I": 60, "J": 40})
setup_print(ws)

# ==========================================================================
# COMPRAVENTA (flujo mensual)
# ==========================================================================
ws = ws_flip
title(ws, "Estrategia 1 — Compraventa (adjudicación y reventa)",
      "Flujo de caja mensual desde el cierre de la subasta. Todas las cifras en euros.")
put(ws, "B4", "Escenario activo:", F_LBL_B)
put(ws, "C4", "=NombreEscenario", F_LINK_B)

N_MONTHS = 24
FIRST_COL = 4  # D
cols = [get_column_letter(FIRST_COL + i) for i in range(N_MONTHS)]
LASTC = cols[-1]
T0 = 27  # fila de cabecera de la tabla mensual
put(ws, f"B{T0-1}", "Flujo de caja mensual", F_SEC)
header(ws, T0, [("B", "Concepto"), ("C", "Total")] + [(c, "") for c in cols])
put(ws, f"B{T0+1}", "Mes", F_LBL_B)
put(ws, f"B{T0+2}", "Fecha", F_LBL_B)
put(ws, f"B{T0+3}", "Hito", F_LBL_B)
for i, c in enumerate(cols):
    if i == 0:
        put(ws, f"{c}{T0+1}", 0, F_IN, fmt=IDX, align=CENTER)
    else:
        put(ws, f"{c}{T0+1}", f"={cols[i-1]}{T0+1}+1", F_FX, fmt=IDX, align=CENTER)
    put(ws, f"{c}{T0+2}", f"=EDATE(FechaCierre,{c}${T0+1})", F_FX, fmt=MONTH, align=CENTER)
    m = f"{c}${T0+1}"
    put(ws, f"{c}{T0+3}",
        f'=TRIM(IF({m}=0,"Subasta ","")&IF({m}=MesPago,"Pago ","")&IF({m}=MesDecreto,"Decreto ","")'
        f'&IF({m}=MesPosesion,"Posesión ","")&IF({m}=MesVenta,"Venta",""))',
        Font(name=FONT, size=8, italic=True, color=GREY), align=CENTER)

flows = [
    ("Depósito (5% del valor de subasta)", "=IF({m}=0,-Deposito,0)"),
    ("Resto del precio de remate", "=IF({m}=MesPago,-(Puja-Deposito),0)"),
    ("Impuesto de adquisición (ITP / IVA)", "=IF({m}=MesDecreto,-ImpuestoAdq,0)"),
    ("Registro, gestoría y abogado", "=IF({m}=MesDecreto,-(Registro+Gestoria+Abogado),0)"),
    ("Deudas afectas (comunidad + IBI)", "=IF({m}=MesPosesion,-(DeudaComunidad+DeudaIBI),0)"),
    ("Toma de posesión", "=IF({m}=MesPosesion,-CostePosesion,0)"),
    ("Gastos corrientes (comunidad, IBI, seguro)", "=IF(AND({m}>MesDecreto,{m}<=MesVenta),-GastoMensual,0)"),
    ("Ingreso neto por la venta", "=IF({m}=MesVenta,VentaNeta,0)"),
]
fr = T0 + 4
FLOW_ROWS = []
for lbl, fx in flows:
    put(ws, f"B{fr}", lbl, F_LBL)
    for c in cols:
        put(ws, f"{c}{fr}", fx.format(m=f"{c}${T0+1}"), F_FX, fmt=EUR0)
    put(ws, f"C{fr}", f"=SUM({cols[0]}{fr}:{LASTC}{fr})", F_FX_B, fmt=EUR0)
    FLOW_ROWS.append(fr)
    fr += 1
R_PRE = fr
put(ws, f"B{R_PRE}", "Flujo antes de impuestos", F_LBL_B)
for c in cols + ["C"]:
    put(ws, f"{c}{R_PRE}", f"=SUM({c}{FLOW_ROWS[0]}:{c}{FLOW_ROWS[-1]})", F_FX_B, fmt=EUR0, border=TOPLINE)
R_TAX = R_PRE + 1
put(ws, f"B{R_TAX}", "Impuesto sobre la ganancia", F_LBL)
for c in cols:
    put(ws, f"{c}{R_TAX}", f"=IF({c}${T0+1}=MesVenta,-MAX(0,$C${R_PRE})*TipoGanancia,0)", F_FX, fmt=EUR0)
put(ws, f"C{R_TAX}", f"=SUM({cols[0]}{R_TAX}:{LASTC}{R_TAX})", F_FX_B, fmt=EUR0)
R_NET = R_TAX + 1
put(ws, f"B{R_NET}", "Flujo después de impuestos", F_LBL_B)
for c in cols + ["C"]:
    put(ws, f"{c}{R_NET}", f"={c}{R_PRE}+{c}{R_TAX}", F_FX_B, fmt=EUR0, border=TOPLINE)
R_CUM = R_NET + 1
put(ws, f"B{R_CUM}", "Caja acumulada", F_LBL)
for i, c in enumerate(cols):
    prev = f"{cols[i-1]}{R_CUM}+" if i else ""
    put(ws, f"{c}{R_CUM}", f"={prev}{c}{R_NET}", F_FX, fmt=EUR0)

# KPIs
put(ws, "B6", "Indicadores", F_SEC)
inv_fx = f"=-SUM(C{FLOW_ROWS[0]}:C{FLOW_ROWS[-2]})"
kpis = [
    ("Puja (precio de remate)", "=Puja", EUR0, F_LINK, None),
    ("Impuesto de adquisición", "=ImpuestoAdq", EUR0, F_LINK, None),
    ("Costes fijos de adquisición", "=CostesFijosAdq", EUR0, F_LINK, None),
    ("Coste de tenencia", "=CosteTenenciaFlip", EUR0, F_LINK, None),
    ("Inversión total", inv_fx, EUR0, F_FX_B, "FlipInversion"),
    ("Ingreso neto por la venta", f"=C{FLOW_ROWS[-1]}", EUR0, F_FX, "FlipVentaNeta"),
    ("Beneficio antes de impuestos", f"=C{R_PRE}", EUR0, F_FX_B, "FlipBAI"),
    ("Impuesto sobre la ganancia", f"=C{R_TAX}", EUR0, F_FX, None),
    ("Beneficio neto", f"=C{R_NET}", EUR0, F_FX_B, "FlipNeto"),
    ("ROI antes de impuestos", f"=IF(C11>0,C13/C11,0)", PCT1, F_FX_B, "FlipROI"),
    ("ROI neto", f"=IF(C11>0,C15/C11,0)", PCT1, F_FX, "FlipROINeto"),
    ("Múltiplo sobre el capital (MOIC neto)", f"=IF(C11>0,(C11+C15)/C11,0)", MULT, F_FX, None),
    ("Plazo total (meses)", "=MesVenta", IDX, F_LINK, None),
    ("TIR anual antes de impuestos", f"=(1+IRR({cols[0]}{R_PRE}:{LASTC}{R_PRE},0.02))^12-1", PCT1, F_FX_B, "FlipTIR"),
    ("TIR anual neta", f"=(1+IRR({cols[0]}{R_NET}:{LASTC}{R_NET},0.02))^12-1", PCT1, F_FX, "FlipTIRNeta"),
    ("Exposición máxima de caja", f"=-MIN({cols[0]}{R_CUM}:{LASTC}{R_CUM})", EUR0, F_FX, None),
]
rr = 7
for lbl, fx, fmt, font, nm in kpis:
    put(ws, f"B{rr}", lbl, F_LBL_B if font in (F_FX_B,) else F_LBL)
    put(ws, f"C{rr}", fx, font, fmt=fmt, border=BOX, fill=FILL_KPI if font is F_FX_B else None)
    if nm:
        add_name(nm, ws, f"C{rr}")
    rr += 1
assert ws["B11"].value == "Inversión total" and ws["B13"].value == "Beneficio antes de impuestos" and ws["B15"].value == "Beneficio neto"
put(ws, f"B{rr}", "Control: fórmula cerrada = flujo mensual", F_LBL)
put(ws, f"C{rr}",
    f'=IF(ABS(C13-(VentaNeta-Puja-ImpuestoAdq-CostesFijosAdq-CosteTenenciaFlip))<0.01,"OK","REVISAR")',
    F_FX_B, border=BOX)
add_name("FlipControl", ws, f"C{rr}")
put(ws, f"B{rr+1}", "Control: venta dentro del horizonte modelizado (mes ≤ 23)", F_LBL)
put(ws, f"C{rr+1}", '=IF(MesVenta<=23,"OK","AMPLIAR HORIZONTE")', F_FX_B, border=BOX)
put(ws, "E7", "Notas", F_LBL_B)
notes = [
    "• El depósito se consigna al pujar; si no se gana la subasta se devuelve íntegro.",
    "• El impuesto de adquisición se liquida tras el decreto de adjudicación (30 días hábiles).",
    "• Las deudas afectas de comunidad e IBI se pagan al tomar posesión para regularizar la plaza.",
    "• El impuesto sobre la ganancia se imputa en el mes de venta (simplificación: se paga en la renta del año siguiente).",
    "• La TIR es mensual anualizada sobre los flujos de la tabla inferior.",
]
for i, n in enumerate(notes):
    put(ws, f"E{8+i}", n, F_NOTE)

widths(ws, {"A": 2, "B": 54, "C": 14})
for c in cols:
    ws.column_dimensions[c].width = 10.5
ws.freeze_panes = f"D{T0+1}"
setup_print(ws)

# ==========================================================================
# ALQUILER (proyección anual a 10 años)
# ==========================================================================
ws = ws_alq
title(ws, "Estrategia 2 — Alquiler (patrimonial a 10 años)",
      "Año 0 = desde el cierre de la subasta hasta el inicio del alquiler. Análisis antes de impuestos.")
put(ws, "B4", "Escenario activo:", F_LBL_B)
put(ws, "C4", "=NombreEscenario", F_LINK_B)
YC = [get_column_letter(4 + i) for i in range(11)]  # D..N años 0..10
A0 = 24
put(ws, f"B{A0-1}", "Proyección anual", F_SEC)
header(ws, A0, [("B", "Concepto"), ("C", "Total")] + [(c, "") for c in YC])
put(ws, f"B{A0+1}", "Año", F_LBL_B)
for i, c in enumerate(YC):
    put(ws, f"{c}{A0+1}", 0 if i == 0 else f"={YC[i-1]}{A0+1}+1", F_IN if i == 0 else F_FX, fmt=IDX, align=CENTER)
y = lambda c: f"{c}${A0+1}"
lines = [
    ("inv", "Inversión inicial (puja + impuestos + costes + tenencia)", lambda c: f"=IF({y(c)}=0,-(Puja+ImpuestoAdq+CostesFijosAdq+CosteTenenciaAlq),0)"),
    ("rent", "Renta de mercado (lo que paga el inquilino)", lambda c: f"=IF({y(c)}>=1,Renta*12*(1+CrecRenta)^({y(c)}-1),0)"),
    ("iva", "IVA repercutido incluido en la renta", lambda c: f"=-{c}{{rent}}*TipoIVA*RentaInclIVA/(1+TipoIVA*RentaInclIVA)"),
    ("vac", "Desocupación e impagos", lambda c: f"=-({c}{{rent}}+{c}{{iva}})*Desocupacion"),
    ("ing", "Ingresos netos", lambda c: f"={c}{{rent}}+{c}{{iva}}+{c}{{vac}}"),
    ("com", "Comunidad", lambda c: f"=IF({y(c)}>=1,-CuotaComunidad*12*(1+InflGastos)^({y(c)}-1),0)"),
    ("ibi", "IBI", lambda c: f"=IF({y(c)}>=1,-IBIAnual*(1+InflGastos)^({y(c)}-1),0)"),
    ("seg", "Seguro", lambda c: f"=IF({y(c)}>=1,-SeguroAnual*(1+InflGastos)^({y(c)}-1),0)"),
    ("ges", "Gestión", lambda c: f"=-{c}{{ing}}*GestionPct"),
    ("man", "Reserva de mantenimiento", lambda c: f"=-{c}{{ing}}*MantenimientoPct"),
    ("noi", "Resultado operativo neto (NOI)", lambda c: f"=SUM({c}{{ing}}:{c}{{man}})-{c}{{ing}}+{c}{{ing}}"),
    ("exit", "Venta al final del año 10 (neta de gastos)", lambda c: f"=IF({y(c)}=10,PrecioVenta*(1+Revalorizacion)^10*(1-DescuentoNeg)*(1-ComisionVenta*(1+TipoIVA))-PlusvaliaSalida,0)"),
    ("cf", "Flujo de caja", lambda c: f"={c}{{inv}}+{c}{{noi}}+{c}{{exit}}"),
    ("cum", "Caja acumulada", None),
]
ROW = {}
rr = A0 + 2
for key, _, _ in lines:
    ROW[key] = rr
    rr += 1
for key, lbl, fn in lines:
    row = ROW[key]
    bold = key in ("ing", "noi", "cf")
    put(ws, f"B{row}", lbl, F_LBL_B if bold else F_LBL)
    for i, c in enumerate(YC):
        if key == "cum":
            fx = f"={YC[i-1]}{row}+{c}{ROW['cf']}" if i else f"={c}{ROW['cf']}"
        elif key == "noi":
            fx = f"={c}{ROW['ing']}+SUM({c}{ROW['com']}:{c}{ROW['man']})"
        else:
            fx = fn(c).format(**ROW)
        put(ws, f"{c}{row}", fx, F_FX_B if bold else F_FX, fmt=EUR0,
            border=TOPLINE if key in ("ing", "noi", "cf") else None)
    if key != "cum":
        put(ws, f"C{row}", f"=SUM({YC[0]}{row}:{YC[-1]}{row})", F_FX_B if bold else F_FX, fmt=EUR0,
            border=TOPLINE if key in ("ing", "noi", "cf") else None)

put(ws, "B6", "Indicadores", F_SEC)
akpis = [
    ("Inversión total", f"=-{YC[0]}{ROW['inv']}", EUR0, F_FX_B, "AlqInversion"),
    ("Renta de mercado (año 1)", f"={YC[1]}{ROW['rent']}", EUR0, F_FX, None),
    ("Ingresos netos (año 1)", f"={YC[1]}{ROW['ing']}", EUR0, F_FX, None),
    ("NOI (año 1)", f"={YC[1]}{ROW['noi']}", EUR0, F_FX_B, "AlqNOI"),
    ("Rentabilidad bruta (renta neta de IVA / inversión)", "=IF(C7>0,RentaNeta*12/C7,0)", PCT1, F_FX, "AlqRentBruta"),
    ("Rentabilidad neta (NOI / inversión)", "=IF(C7>0,C10/C7,0)", PCT1, F_FX_B, "AlqRentNeta"),
    ("Rentabilidad bruta de anuncio (renta × 12 / puja)", "=IF(Puja>0,Renta*12/Puja,0)", PCT1, F_FX, None),
    ("TIR a 10 años (con venta final)", f"=IRR({YC[0]}{ROW['cf']}:{YC[-1]}{ROW['cf']},0.05)", PCT1, F_FX_B, "AlqTIR"),
    ("VAN a la tasa de descuento", f"={YC[0]}{ROW['cf']}+NPV(TasaDescuento,{YC[1]}{ROW['cf']}:{YC[-1]}{ROW['cf']})", EUR0, F_FX, "AlqVAN"),
    ("Años para recuperar la inversión (simple)", '=IF(C10>0,C7/C10,"n/a")', NUM2, F_FX, None),
    ("Múltiplo sobre el capital (MOIC)", f"=IF(C7>0,(C7+C{ROW['cf']})/C7,0)", MULT, F_FX, None),
]
rr = 7
for lbl, fx, fmt, font, nm in akpis:
    put(ws, f"B{rr}", lbl, F_LBL_B if font is F_FX_B else F_LBL)
    put(ws, f"C{rr}", fx, font, fmt=fmt, border=BOX, fill=FILL_KPI if font is F_FX_B else None)
    if nm:
        add_name(nm, ws, f"C{rr}")
    rr += 1
assert ws["B7"].value == "Inversión total" and ws["B10"].value == "NOI (año 1)"
put(ws, f"B{rr}", "Control: NOI año 1 = fórmula cerrada", F_LBL)
put(ws, f"C{rr}", '=IF(ABS(C10-NOIInicial)<0.01,"OK","REVISAR")', F_FX_B, border=BOX)
add_name("AlqControl", ws, f"C{rr}")
put(ws, "E7", "Notas", F_LBL_B)
for i, n in enumerate([
    "• Alquilar una plaza de garaje independiente obliga a repercutir IVA (21%): la renta de mercado se trata como IVA incluido.",
    "• Si el inversor es empresario a efectos de IVA, puede deducir el IVA soportado en la compra (ver Supuestos, sección 4).",
    "• Análisis antes de impuestos sobre la renta (IRPF / IS) y sin amortización fiscal.",
    "• El valor de salida aplica la revalorización anual al precio de salida actual y descuenta comisión y plusvalía.",
]):
    put(ws, f"E{8+i}", n, F_NOTE)
widths(ws, {"A": 2, "B": 52, "C": 14})
for c in YC:
    ws.column_dimensions[c].width = 11
ws.freeze_panes = f"D{A0+1}"
setup_print(ws)

# ==========================================================================
# PUJA MÁXIMA
# ==========================================================================
ws = ws_pm
title(ws, "Puja máxima y escalera de pujas",
      "Importe máximo que se puede ofrecer para alcanzar la rentabilidad objetivo, y resultado de cada nivel de puja.")
section(ws, 5, "A. Puja máxima por estrategia", "B", "L")
header(ws, 6, [("B", "Concepto"), ("C", "Compraventa"), ("D", "Alquiler")])
pm_rows = [
    ("Rentabilidad objetivo", "=ROIObjetivo", "=RentabObjetivo", PCT1, F_LINK),
    ("Ingreso neto por la venta / NOI año 1", "=VentaNeta", "=NOIInicial", EUR0, F_LINK),
    ("Inversión total máxima admisible", "=C8/(1+C7)", "=IF(D7>0,D8/D7,0)", EUR0, F_FX),
    ("(−) Costes fijos de adquisición", "=-CostesFijosAdq", "=-CostesFijosAdq", EUR0, F_LINK),
    ("(−) Coste de tenencia", "=-CosteTenenciaFlip", "=-CosteTenenciaAlq", EUR0, F_LINK),
    ("Disponible para puja + impuesto", "=SUM(C9:C11)", "=SUM(D9:D11)", EUR0, F_FX),
    ("Puja máxima", "=" + bid_from_avail("C12"), "=" + bid_from_avail("D12"), EUR0, F_FX_B),
    ("Puja máxima / valor de subasta", "=C13/ValorSubasta", "=D13/ValorSubasta", PCT1, F_FX),
    ("¿Supera el umbral del 70%?", '=IF(C13>=UmbralSetenta,"Sí","No — riesgo art. 670.4 LEC")',
     '=IF(D13>=UmbralSetenta,"Sí","No — riesgo art. 670.4 LEC")', None, F_FX),
    ("Puja de equilibrio (beneficio / rentabilidad 0)", "=" + bid_from_avail("C8-CostesFijosAdq-CosteTenenciaFlip"), "", EUR0, F_FX),
]
rr = 7
for lbl, a, b, fmt, font in pm_rows:
    put(ws, f"B{rr}", lbl, F_LBL_B if font is F_FX_B else F_LBL)
    put(ws, f"C{rr}", a, font, fmt=fmt, border=BOX, fill=FILL_KPI if font is F_FX_B else None)
    if b != "":
        put(ws, f"D{rr}", b, font, fmt=fmt, border=BOX, fill=FILL_KPI if font is F_FX_B else None)
    rr += 1
add_name("PujaMaxFlip", ws, "C13")
add_name("PujaMaxAlq", ws, "D13")
add_name("PujaEquilibrio", ws, "C16")
put(ws, "B18", "Puja máxima recomendada (estrategia principal)", F_LBL_B)
put(ws, "C18", "=IF(Estrategia=1,PujaMaxFlip,PujaMaxAlq)", F_FX_B, fmt=EUR0, border=BOX, fill=FILL_KEY)
add_name("PujaMaxSel", ws, "C18")
put(ws, "B19", "Puja mínima prudente (umbral 70%, redondeado al euro superior)", F_LBL)
put(ws, "C19", "=ROUNDUP(UmbralSetenta,0)", F_FX, fmt=EUR0, border=BOX)
add_name("PujaMinPrudente", ws, "C19")
put(ws, "E7", "Cómo se calcula", F_LBL_B)
for i, n in enumerate([
    "• Compraventa: inversión máxima = ingreso neto por venta / (1 + ROI objetivo).",
    "• Alquiler: inversión máxima = NOI año 1 / rentabilidad neta objetivo.",
    "• Se restan los costes fijos y de tenencia; el resto se reparte entre puja e impuesto.",
    "• Con ITP la base es el mayor entre la puja y el valor de referencia catastral.",
    "• Con IVA no deducible: puja máxima = disponible / (1 + 21%).",
]):
    put(ws, f"E{8+i}", n, F_NOTE)

section(ws, 21, "B. Escalera de pujas", "B", "L")
header(ws, 22, [("B", "% valor subasta"), ("C", "Puja"), ("D", "Régimen art. 670 LEC"), ("E", "Impuesto adq."),
                ("F", "Inversión compraventa"), ("G", "Beneficio antes imp."), ("H", "ROI compraventa"),
                ("I", "ROI anualizado"), ("J", "Beneficio neto"), ("K", "Rentab. neta alquiler"), ("L", "¿Cumple objetivo?")])
L0 = 23
pcts = [round(0.40 + 0.05 * i, 2) for i in range(23)]  # 40% .. 150%
for i, p in enumerate(pcts):
    rr = L0 + i
    put(ws, f"B{rr}", p, F_IN, fmt='0%', align=CENTER)
    put(ws, f"C{rr}", f"=B{rr}*ValorSubasta", F_FX, fmt=EUR0)
    put(ws, f"D{rr}", f'=IF(C{rr}>=UmbralSetenta,"Aprobación directa (≥ 70%)",IF(C{rr}>=UmbralCincuenta,"Mejorable por ejecutado/ejecutante","< 50%: decide el LAJ"))',
        Font(name=FONT, size=9, color=BLACK))
    put(ws, f"E{rr}", "=" + tax_expr(f"C{rr}"), F_FX, fmt=EUR0)
    put(ws, f"F{rr}", f"=C{rr}+E{rr}+CostesFijosAdq+CosteTenenciaFlip", F_FX, fmt=EUR0)
    put(ws, f"G{rr}", f"=VentaNeta-F{rr}", F_FX, fmt=EUR0)
    put(ws, f"H{rr}", f"=IF(F{rr}>0,G{rr}/F{rr},0)", F_FX, fmt=PCT1)
    put(ws, f"I{rr}", f"=IF(MesVenta>0,(1+H{rr})^(12/MesVenta)-1,0)", F_FX, fmt=PCT1)
    put(ws, f"J{rr}", f"=G{rr}-MAX(0,G{rr})*TipoGanancia", F_FX, fmt=EUR0)
    put(ws, f"K{rr}", f"=NOIInicial/(C{rr}+E{rr}+CostesFijosAdq+CosteTenenciaAlq)", F_FX, fmt=PCT1)
    put(ws, f"L{rr}", f'=IF(AND(C{rr}>=UmbralSetenta,IF(Estrategia=1,H{rr}>=ROIObjetivo,K{rr}>=RentabObjetivo)),"Sí","No")', F_FX, align=CENTER)
L1 = L0 + len(pcts) - 1
ws.conditional_formatting.add(f"L{L0}:L{L1}", CellIsRule(operator="equal", formula=['"Sí"'], fill=FILL_OK))
ws.conditional_formatting.add(f"L{L0}:L{L1}", CellIsRule(operator="equal", formula=['"No"'], fill=FILL_BAD))
ws.conditional_formatting.add(f"G{L0}:G{L1}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, color="C00000")))
ws.conditional_formatting.add(f"D{L0}:D{L1}", FormulaRule(formula=[f"$C{L0}<UmbralSetenta"], fill=FILL_WARN))

ch = LineChart()
ch.title = "Rentabilidad según la puja"
ch.y_axis.title = "Rentabilidad"
ch.x_axis.title = "Puja (€)"
ch.y_axis.number_format = '0%'
ch.y_axis.majorGridlines = None
ch.height, ch.width = 10, 24
C0, C1 = L0 + 2, L0 + 18  # pujas del 50% al 130% del valor de subasta
for col in (8, 11):       # ROI compraventa · rentabilidad neta alquiler
    ser = Reference(ws, min_col=col, min_row=C0, max_row=C1)
    ch.add_data(ser, titles_from_data=False)
from openpyxl.chart.series import SeriesLabel
ch.series[0].tx = SeriesLabel(v="ROI compraventa")
ch.series[1].tx = SeriesLabel(v="Rentabilidad neta alquiler")
ch.set_categories(Reference(ws, min_col=3, min_row=C0, max_row=C1))
ch.x_axis.tickLblSkip = 2
ch.y_axis.scaling.min = -0.4
ch.y_axis.scaling.max = 1.0
ch.y_axis.majorUnit = 0.2
ch.x_axis.number_format = '#,##0'
for s, color in zip(ch.series, ("2A78D6", "EB6834")):
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = 25400  # 2 pt
    s.smooth = False
ch.legend.position = "b"
ch.x_axis.delete = False
ch.y_axis.delete = False
ws.add_chart(ch, f"B{L1+3}")
widths(ws, {"A": 2, "B": 50, "C": 14, "D": 32, "E": 13, "F": 15, "G": 15, "H": 13, "I": 13, "J": 14, "K": 15, "L": 13})
ws.freeze_panes = "A7"
setup_print(ws)

# ==========================================================================
# SENSIBILIDAD
# ==========================================================================
ws = ws_sen
title(ws, "Análisis de sensibilidad",
      "Cada celda recalcula el resultado con fórmula cerrada; el resto de supuestos son los del escenario activo. Encabezados editables (azul).")
bids = [4000, 4500, 5000, 5500, 6000, 6500, 7000, 7500, 8000]
prices = [7000, 8000, 9000, 10000, 11000, 12000, 13000]
rents = [30, 35, 40, 45, 50, 55, 60]
GC = ["C", "D", "E", "F", "G", "H", "I"]


def grid(r0, ttl, col_hdr_label, col_vals, cell_fx, fmt, scale):
    section(ws, r0, ttl, "B", "I")
    put(ws, f"B{r0+1}", f"Puja ↓ / {col_hdr_label} →", F_HDR, fill=FILL_HDR, align=CENTER)
    for c, v in zip(GC, col_vals):
        put(ws, f"{c}{r0+1}", v, Font(name=FONT, size=10, bold=True, color=BLUE), fmt=EUR0, fill=FILL_SEC, align=CENTER, border=BOX)
    for i, b in enumerate(bids):
        rr = r0 + 2 + i
        put(ws, f"B{rr}", b, Font(name=FONT, size=10, bold=True, color=BLUE), fmt=EUR0, fill=FILL_SEC, align=CENTER, border=BOX)
        for c in GC:
            put(ws, f"{c}{rr}", "=" + cell_fx(f"$B{rr}", f"{c}${r0+1}"), F_FX, fmt=fmt, border=BOX, align=CENTER)
    rng = f"C{r0+2}:I{r0+1+len(bids)}"
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="F4B6B5",
                                                      mid_type="num", mid_value=scale, mid_color="F2F2F2",
                                                      end_type="max", end_color="9FD5B8"))
    return r0 + 2 + len(bids) + 1


def flip_profit(b, p):
    return f"({sale_net_of(p)})-({b}+{tax_expr(b)}+CostesFijosAdq+CosteTenenciaFlip)"


def flip_roi(b, p):
    return f"({flip_profit(b, p)})/({b}+{tax_expr(b)}+CostesFijosAdq+CosteTenenciaFlip)"


def rent_yield(b, rent):
    return f"({noi_of(rent)})/({b}+{tax_expr(b)}+CostesFijosAdq+CosteTenenciaAlq)"


put(ws, "B5", "Caso activo — puja:", F_LBL_B)
put(ws, "C5", "=Puja", F_LINK, fmt=EUR0)
put(ws, "D5", "precio de salida:", F_LBL_B)
put(ws, "E5", "=PrecioVenta", F_LINK, fmt=EUR0)
put(ws, "F5", "renta:", F_LBL_B)
put(ws, "G5", "=Renta", F_LINK, fmt=EUR0)
nr = grid(7, "1. Compraventa — beneficio antes de impuestos (€): puja × precio de salida", "Precio de salida", prices, flip_profit, EUR0, 0)
nr = grid(nr + 1, "2. Compraventa — ROI sobre inversión: puja × precio de salida", "Precio de salida", prices, flip_roi, PCT1, 0)
nr = grid(nr + 1, "3. Alquiler — rentabilidad neta (NOI / inversión): puja × renta mensual", "Renta mensual", rents, rent_yield, PCT1, 0.05)
put(ws, f"B{nr+1}", "Escala de color: rojo = peor, gris = umbral (0 € / 0% en compraventa; 5% en alquiler), verde = mejor.", F_NOTE)
widths(ws, {"A": 2, "B": 26, "C": 13, "D": 13, "E": 13, "F": 13, "G": 13, "H": 13, "I": 13})
setup_print(ws)

# ==========================================================================
# ESCENARIOS
# ==========================================================================
ws = ws_esc
title(ws, "Escenarios — pesimista, base y optimista",
      "Resultados calculados en paralelo para los tres escenarios de la hoja Supuestos (tabla 2).")
header(ws, 5, [("B", "Concepto"), ("C", "Pesimista"), ("D", "Base"), ("E", "Optimista")])
SCOL = {"C": "C", "D": "D", "E": "E"}
erows = [
    ("Escenario activo", lambda c, i: f'=IF(Escenario={i},"◄ activo","")', None, F_FX),
    ("VARIABLES", None, None, None),
    ("Puja", lambda c, i: f"=Supuestos!{c}{SCN['Puja']}", EUR0, F_LINK),
    ("Puja / valor de subasta", lambda c, i: f"={c}8/ValorSubasta", PCT1, F_FX),
    ("Precio de salida", lambda c, i: f"=Supuestos!{c}{SCN['PrecioVenta']}", EUR0, F_LINK),
    ("Renta mensual", lambda c, i: f"=Supuestos!{c}{SCN['Renta']}", EUR0, F_LINK),
    ("Mes de toma de posesión", lambda c, i: f"=Supuestos!{c}{SCN['MesPosesion']}", IDX, F_LINK),
    ("Coste de toma de posesión", lambda c, i: f"=Supuestos!{c}{SCN['CostePosesion']}", EUR0, F_LINK),
    ("Cuota de comunidad (€/mes)", lambda c, i: f"=Supuestos!{c}{SCN['CuotaComunidad']}", EUR0, F_LINK),
    ("Meses de cuotas impagadas afectas", lambda c, i: f"=Supuestos!{c}{SCN['MesesDeuda']}", IDX, F_LINK),
    ("Meses de comercialización", lambda c, i: f"=Supuestos!{c}{SCN['MesesVenta']}", IDX, F_LINK),
    ("Régimen fiscal (1 = ITP · 2 = IVA)", lambda c, i: f"=Supuestos!{c}{SCN['Regimen']}", IDX, F_LINK),
    ("COMPRAVENTA", None, None, None),
    ("Impuesto de adquisición", lambda c, i: "=" + tax_expr(f"{c}8", f"{c}17"), EUR0, F_FX),
    ("Costes fijos de adquisición", lambda c, i: f"=Registro+Gestoria+Abogado+{c}14*{c}15+DeudaIBI+{c}13", EUR0, F_FX),
    ("Mes de venta", lambda c, i: f"={c}12+{c}16", IDX, F_FX),
    ("Coste de tenencia", lambda c, i: f"=({c}14+(IBIAnual+SeguroAnual)/12)*MAX(0,{c}21-MesDecreto)", EUR0, F_FX),
    ("Inversión total", lambda c, i: f"={c}8+{c}19+{c}20+{c}22", EUR0, F_FX_B),
    ("Ingreso neto por la venta", lambda c, i: "=" + sale_net_of(f"{c}10"), EUR0, F_FX),
    ("Beneficio antes de impuestos", lambda c, i: f"={c}24-{c}23", EUR0, F_FX_B),
    ("Beneficio neto", lambda c, i: f"={c}25-MAX(0,{c}25)*TipoGanancia", EUR0, F_FX),
    ("ROI antes de impuestos", lambda c, i: f"=IF({c}23>0,{c}25/{c}23,0)", PCT1, F_FX_B),
    ("ROI anualizado", lambda c, i: f"=IF({c}21>0,(1+{c}27)^(12/{c}21)-1,0)", PCT1, F_FX),
    ("Puja máxima (ROI objetivo)", lambda c, i: "=" + bid_from_avail(f"{c}24/(1+ROIObjetivo)-{c}20-{c}22", f"{c}17"), EUR0, F_FX_B),
    ("ALQUILER", None, None, None),
    ("Coste de tenencia hasta alquilar", lambda c, i: f"=({c}14+(IBIAnual+SeguroAnual)/12)*MAX(0,{c}12+MesesPuestaAlquiler-MesDecreto)", EUR0, F_FX),
    ("Inversión total", lambda c, i: f"={c}8+{c}19+{c}20+{c}31", EUR0, F_FX_B),
    ("NOI año 1", lambda c, i: "=" + noi_of(f"{c}11", f"{c}14"), EUR0, F_FX),
    ("Rentabilidad neta", lambda c, i: f"=IF({c}32>0,{c}33/{c}32,0)", PCT1, F_FX_B),
    ("Puja máxima (rentabilidad objetivo)", lambda c, i: "=" + bid_from_avail(f"{c}33/RentabObjetivo-{c}20-{c}31", f"{c}17"), EUR0, F_FX_B),
]
rr = 6
for lbl, fn, fmt, font in erows:
    if fn is None:
        section(ws, rr, lbl, "B", "E")
        rr += 1
        continue
    put(ws, f"B{rr}", lbl, F_LBL_B if font is F_FX_B else F_LBL)
    for i, c in enumerate(("C", "D", "E"), start=1):
        put(ws, f"{c}{rr}", fn(c, i), font, fmt=fmt, border=BOX, align=CENTER if fmt is None else None)
    rr += 1
# controles de posición de las filas referenciadas en las fórmulas
expected = {8: "Puja", 10: "Precio de salida", 11: "Renta mensual", 12: "Mes de toma de posesión", 13: "Coste de toma de posesión",
            14: "Cuota de comunidad (€/mes)", 15: "Meses de cuotas impagadas afectas", 16: "Meses de comercialización",
            17: "Régimen fiscal (1 = ITP · 2 = IVA)", 19: "Impuesto de adquisición", 20: "Costes fijos de adquisición",
            21: "Mes de venta", 22: "Coste de tenencia", 23: "Inversión total", 24: "Ingreso neto por la venta",
            25: "Beneficio antes de impuestos", 27: "ROI antes de impuestos", 31: "Coste de tenencia hasta alquilar",
            32: "Inversión total", 33: "NOI año 1"}
for k, v in expected.items():
    assert ws[f"B{k}"].value == v, (k, ws[f"B{k}"].value, v)
for nm, cell in (("EscBAI_P", "C25"), ("EscBAI_B", "D25"), ("EscBAI_O", "E25"), ("EscROI_P", "C27"), ("EscROI_B", "D27"),
                 ("EscROI_O", "E27"), ("EscRN_P", "C34"), ("EscRN_B", "D34"), ("EscRN_O", "E34"),
                 ("EscPM_P", "C29"), ("EscPM_B", "D29"), ("EscPM_O", "E29")):
    add_name(nm, ws, cell)
ws.conditional_formatting.add("C6:E6", CellIsRule(operator="equal", formula=['"◄ activo"'], fill=FILL_KEY))
ws.conditional_formatting.add("C25:E26", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, bold=True, color="C00000")))

rr += 1
section(ws, rr, "Contraste fiscal en el escenario activo: ITP frente a IVA", "B", "E")
rr += 1
header(ws, rr, [("B", "Concepto"), ("C", "ITP 7%"), ("D", "IVA 21%"), ("E", "Diferencia")])
rr += 1
put(ws, f"B{rr}", "Impuesto de adquisición", F_LBL)
put(ws, f"C{rr}", "=TipoITP*MAX(Puja,ValorReferencia)", F_FX, fmt=EUR0, border=BOX)
put(ws, f"D{rr}", "=TipoIVA*Puja*(1-IVADeducible)", F_FX, fmt=EUR0, border=BOX)
put(ws, f"E{rr}", f"=D{rr}-C{rr}", F_FX_B, fmt=EUR0, border=BOX)
add_name("ImpactoIVA", ws, f"E{rr}")
rr += 1
put(ws, f"B{rr}", "Beneficio compraventa antes de impuestos", F_LBL)
put(ws, f"C{rr}", f"=VentaNeta-(Puja+C{rr-1}+CostesFijosAdq+CosteTenenciaFlip)", F_FX, fmt=EUR0, border=BOX)
put(ws, f"D{rr}", f"=VentaNeta-(Puja+D{rr-1}+CostesFijosAdq+CosteTenenciaFlip)", F_FX, fmt=EUR0, border=BOX)
put(ws, f"E{rr}", f"=D{rr}-C{rr}", F_FX_B, fmt=EUR0, border=BOX)
widths(ws, {"A": 2, "B": 46, "C": 16, "D": 16, "E": 16})
setup_print(ws, landscape=False)

# ==========================================================================
# RIESGOS Y DUE DILIGENCE
# ==========================================================================
ws = ws_rsk
title(ws, "Riesgos, due diligence y marco legal",
      "Riesgos identificados en los documentos del expediente y cómo los recoge el modelo.")
header(ws, 5, [("B", "Riesgo"), ("C", "Descripción"), ("D", "Probabilidad"), ("E", "Impacto"), ("F", "Mitigación"), ("G", "En el modelo")])
risks = [
    ("Régimen fiscal: IVA en lugar de ITP",
     "SUPERCHINDO, S.L. es titular por división horizontal (2018). Si la plaza nunca se ha vendido, la adjudicación puede ser una 'primera entrega' sujeta a IVA 21% en lugar de ITP 7%.",
     "Media", "Medio", "Consultar con un asesor fiscal antes de pujar; preguntar al Juzgado si el remate se entiende con IVA incluido.",
     "Selector de régimen; el escenario pesimista usa IVA. Hoja Escenarios: contraste ITP/IVA."),
    ("Puja inferior al 70% del valor",
     "Con una puja inferior a 4.954,92 €, el ejecutado puede presentar un tercero que mejore la puja y la Comunidad ejecutante puede adjudicarse la plaza por el 70% o por lo que se le debe (10.083,73 €).",
     "Alta (si < 70%)", "Alto", "Pujar al menos el 70% del valor de subasta.",
     "Hoja Puja máxima: régimen art. 670 por nivel de puja."),
    ("Plaza solo apta para moto",
     "Superficie útil registral de 3,75 m² (9,50 m² construidos). La demanda se limita a motos y scooters.",
     "Alta", "Medio", "Verificar medidas reales y acceso (rampa o plataforma) con una visita al garaje por C/ Constancia 71.",
     "Comparables solo de plazas de moto."),
    ("Deudas de comunidad (afección real)",
     "El adjudicatario responde de las cuotas impagadas del año en curso y de los tres años anteriores (art. 9.1.e LPH). El ejecutado ya debe a la Comunidad.",
     "Alta", "Bajo", "Pedir al administrador de la Comunidad un certificado de deuda de la plaza nº 10.",
     "Cuota × meses impagados (Supuestos, tabla 2)."),
    ("IBI pendiente",
     "El inmueble queda afecto al IBI impagado de hasta 4 años (art. 64 TRLRHL).",
     "Media", "Bajo", "Consultar la deuda en la Agencia Tributaria de Sevilla.", "IBI anual × años pendientes."),
    ("Situación posesoria desconocida",
     "El edicto indica que no consta la situación posesoria de la plaza (podría usarla un tercero).",
     "Baja", "Medio", "Pedir al Juzgado la inspección (art. 669.3 LEC); posesión por art. 675 LEC si está ocupada.",
     "Escenario pesimista: posesión en el mes 6 y 800 € de coste."),
    ("Competencia en la puja",
     "Las subastas de garajes de importe bajo atraen a muchos postores; la subasta se prorroga 1 hora tras cada puja (hasta 24 h).",
     "Alta", "Medio", "Fijar la puja máxima de antemano y no superarla.", "Hoja Puja máxima."),
    ("Anuncio de una plaza idéntica en el mismo edificio",
     "Idealista 112744803: plaza de moto en C/ Virtud 44 con 9,50 m² construidos, 10.000 € y posesión 'a partir de noviembre'. O es esta misma finca anunciada antes de la adjudicación, o es una plaza gemela que competirá en la reventa.",
     "Media", "Medio", "Tratarlo como precio objetivo de salida, no como comparable independiente.",
     "Precio de salida base 10.000 € con un 5% de negociación."),
    ("Liquidez de la reventa",
     "El mercado de plazas de moto es estrecho y los plazos de venta pueden alargarse.",
     "Media", "Medio", "Precio competitivo; alquiler como plan B.", "Escenario pesimista: 8.000 € y 8 meses."),
    ("Certificación de cargas de 2023",
     "La certificación es de 24/03/2023. Las cargas posteriores al embargo se cancelan con la adjudicación (art. 674 LEC), pero conviene confirmar la información registral actualizada.",
     "Baja", "Bajo", "Pedir una nota simple de la finca 37.841 (Registro nº 2) o revisar la información registral del Portal (art. 667.2 LEC).",
     "Cargas anteriores = 0 € (Diligencia de 19/03/2025)."),
    ("Cambios societarios del ejecutado",
     "Un directorio empresarial (informa.es) recoge un anuncio de fusión por absorción en el BORME de 11/11/2025 que afecta a SUPERCHINDO, S.L.",
     "Baja", "Medio", "Confirmar con el Juzgado que la sucesión procesal no afecta a la inscripción del decreto.", "No cuantificado."),
    ("Plazos judiciales",
     "La aprobación del remate, el decreto de adjudicación y el testimonio pueden tardar meses.",
     "Media", "Bajo", "Asumir de 3 a 6 meses.", "Calendario por escenario."),
    ("Suspensión de la subasta",
     "El ejecutado puede pagar la deuda antes de aprobarse el remate, o un concurso puede suspender la ejecución.",
     "Baja", "Bajo", "Se devuelve el depósito.", "No cuantificado."),
]
rr = 6
for rk in risks:
    for col, v in zip("BCDEFG", rk):
        f = F_LBL_B if col == "B" else F_LBL
        put(ws, f"{col}{rr}", v, f, align=WRAP, border=BOX)
    for col in "DE":
        v = ws[f"{col}{rr}"].value
        ws[f"{col}{rr}"].fill = FILL_BAD if v.startswith("Alt") else FILL_WARN if v.startswith("Med") else FILL_OK
    ws.row_dimensions[rr].height = 64
    rr += 1

rr += 1
section(ws, rr, "Checklist de due diligence antes de pujar", "B", "G")
rr += 1
header(ws, rr, [("B", "Comprobación"), ("C", "Dónde / cómo"), ("D", "Estado"), ("E", "Responsable"), ("F", "Notas"), ("G", "")])
rr += 1
checks = [
    ("Fechas, tramos y pujas actuales", "Ficha del Portal de Subastas (pestañas Información general y Pujas)."),
    ("Información registral actualizada", "Portal de Subastas (art. 667.2 LEC) o nota simple de la finca 37.841 en el Registro nº 2 de Sevilla."),
    ("Referencia catastral y valor de referencia", "Sede Electrónica del Catastro (C/ Constancia 71, plaza 10)."),
    ("Certificado de deuda de comunidad", "Administrador de la CP Calle Virtud 44."),
    ("Deuda de IBI", "Agencia Tributaria de Sevilla."),
    ("Visita y medidas reales de la plaza", "Solicitud de inspección al JPI nº 4 (art. 669.3 LEC); acceso por C/ Constancia 71."),
    ("Régimen fiscal (ITP / IVA)", "Asesor fiscal; consulta al Juzgado sobre el IVA en el remate."),
    ("Subastas relacionadas (local comercial)", "Pestaña 'Relacionados' de la ficha."),
    ("Alta en el Portal y depósito", "Certificado digital; consignar 353,92 € antes del cierre."),
]
dv_state = DataValidation(type="list", formula1='"Pendiente,En curso,Hecho"', allow_blank=True)
ws.add_data_validation(dv_state)
for chk, how in checks:
    put(ws, f"B{rr}", chk, F_LBL, align=WRAP, border=BOX)
    put(ws, f"C{rr}", how, F_LBL, align=WRAP, border=BOX)
    put(ws, f"D{rr}", "Pendiente", F_IN, border=BOX, align=CENTER)
    dv_state.add(f"D{rr}")
    put(ws, f"E{rr}", "", F_IN, border=BOX)
    put(ws, f"F{rr}", "", F_IN, border=BOX)
    ws.row_dimensions[rr].height = 30
    rr += 1
ws.conditional_formatting.add(f"D{rr-len(checks)}:D{rr-1}", CellIsRule(operator="equal", formula=['"Hecho"'], fill=FILL_OK))
ws.conditional_formatting.add(f"D{rr-len(checks)}:D{rr-1}", CellIsRule(operator="equal", formula=['"Pendiente"'], fill=FILL_WARN))

rr += 1
section(ws, rr, "Marco legal aplicable (resumen)", "B", "G")
rr += 1
legal = [
    ("Art. 647 y 669.1 LEC", "Para pujar hay que consignar el 5% del valor de subasta. El ejecutante puja sin depósito y solo él puede ceder el remate."),
    ("Art. 649 LEC", "La subasta dura 20 días naturales y no se cierra hasta pasada 1 hora desde la última puja (máximo 24 h de ampliación)."),
    ("Art. 652–653 LEC", "Reserva de postura: si el rematante no paga, el remate pasa a los siguientes postores y quien no paga pierde el depósito."),
    ("Art. 670 LEC", "Puja ≥ 70%: se aprueba el remate y se paga el resto en 40 días. Puja < 70%: el ejecutado tiene 10 días para presentar un tercero que mejore la puja, después el ejecutante tiene 5 días para adjudicarse la plaza por el 70% o por lo que se le debe; si no lo hace, se aprueba la puja si es ≥ 50%, y por debajo decide el LAJ."),
    ("Art. 668.2 y 642 LEC", "Las cargas anteriores subsisten y el adjudicatario se subroga en ellas. En esta finca son 0 €."),
    ("Art. 674–675 LEC", "Testimonio del decreto como título inscribible, cancelación de las cargas posteriores y entrega de la posesión."),
    ("Art. 9.1.e LPH", "Afección del inmueble a las cuotas de comunidad del año en curso y los tres anteriores."),
    ("Art. 64 TRLRHL", "Afección del inmueble al IBI impagado."),
    ("ITP Andalucía", "Tipo general del 7% sobre el mayor entre el precio y el valor de referencia del Catastro (Ley 11/2021). El TSJ de Galicia (23/04/2026) aplicó el valor de referencia también en una subasta."),
]
for art, txt in legal:
    put(ws, f"B{rr}", art, F_LBL_B, align=WRAP)
    put(ws, f"C{rr}", txt, F_LBL, align=WRAP)
    ws.merge_cells(f"C{rr}:G{rr}")
    ws.row_dimensions[rr].height = 44 if len(txt) > 160 else 30
    rr += 1
widths(ws, {"A": 2, "B": 30, "C": 62, "D": 13, "E": 12, "F": 46, "G": 36})
setup_print(ws)

# ==========================================================================
# FUENTES
# ==========================================================================
ws = ws_src
title(ws, "Fuentes", "Documentos del expediente y fuentes externas consultadas el 13/10/2026.")
header(ws, 5, [("B", "Fuente"), ("C", "Uso en el modelo"), ("D", "Referencia / enlace")])
srcs = [
    ("Doc. 1 — Edicto de subasta, JPI nº 4 de Sevilla (16/06/2025)", "Descripción del bien, cantidad reclamada, condiciones generales y particulares, situación posesoria.", "Ejecución de títulos judiciales 1455/2022"),
    ("Doc. 2 — Mandamiento y certificación de dominio y cargas, Registro de la Propiedad nº 2 de Sevilla (24/03/2023)", "Titularidad, superficies, cuota, cargas, ausencia de titulares posteriores.", "Finca 37.841 · IDUFIR 41028000513716"),
    ("Doc. 4 — Decreto de valoración definitiva (19/03/2025)", "Valoración pericial de 7.078,45 €.", "Art. 639.4 LEC"),
    ("Doc. 5 — Diligencia de ordenación (19/03/2025)", "Valor a efectos de subasta y cargas deducibles (0 €).", "Art. 666 LEC"),
    ("Portal de Subastas del BOE", "Identificador de la subasta.", "https://subastas.boe.es/detalleSubasta.php?idSub=SUB-JA-2026-265709"),
    ("idealista — comparables de venta y alquiler de plazas de moto", "Precio de salida y renta (hoja Comparables).", "https://www.idealista.com/venta-garajes/sevilla-sevilla/con-moto/"),
    ("RealAdvisor — Calle Virtud (41010)", "Código postal y contexto de precios del barrio.", "https://realadvisor.es/es/precios-viviendas/41010-sevilla/calle-virtud"),
    ("Ley de Enjuiciamiento Civil (texto consolidado)", "Arts. 642, 647, 649, 652–653, 666–675.", "https://www.boe.es/buscar/act.php?id=BOE-A-2000-323"),
    ("Ley de Propiedad Horizontal", "Art. 9.1.e (afección por cuotas de comunidad).", "https://www.boe.es/buscar/act.php?id=BOE-A-1960-10906"),
    ("TSJ de Galicia, 23/04/2026 — valor de referencia en subastas", "Base imponible del ITP en adjudicaciones en subasta.",
     "https://www.poderjudicial.es/cgpj/es/Poder-Judicial/Tribunales-Superiores-de-Justicia/TSJ-Galicia/Noticias-Judiciales-TSJ-Galicia/El-TSXG-ratifica-que-el-valor-de-referencia-de-los-inmuebles-prevalece-sobre-el-de-adquisicion-en-subastas-para-el-calculo-de-la-base-imponible-del-Impuesto-sobre-Transmisiones-Patrimoniales-"),
    ("ITP en Andalucía 2026", "Tipo general del 7%.", "https://taxdown.es/informacion-impuestos/itp-andalucia"),
    ("Directorio informa.es — SUPERCHINDO", "Anuncio de fusión por absorción en el BORME (11/11/2025).", "https://www.informa.es/directorio-empresas/Empresa_SUPERCHINDO.html"),
]
rr = 6
for a, b, c in srcs:
    put(ws, f"B{rr}", a, F_LBL, align=WRAP, border=BOX)
    put(ws, f"C{rr}", b, F_LBL, align=WRAP, border=BOX)
    cell = put(ws, f"D{rr}", c, F_LBL, align=WRAP, border=BOX)
    if c.startswith("http"):
        cell.hyperlink = c
        cell.font = Font(name=FONT, size=9, color="0563C1", underline="single")
    ws.row_dimensions[rr].height = 32
    rr += 1
rr += 1
put(ws, f"B{rr}", ("Limitaciones: la ficha del Portal (fechas, tramos y pujas) no se pudo consultar desde el entorno de trabajo y el documento 3 "
                   "no se aportó. Los comparables son precios de oferta, no de cierre. Este modelo no constituye asesoramiento jurídico ni fiscal."), F_NOTE, align=WRAP)
ws.merge_cells(f"B{rr}:D{rr}")
ws.row_dimensions[rr].height = 40
widths(ws, {"A": 2, "B": 55, "C": 55, "D": 70})
setup_print(ws)

# ==========================================================================
# RESUMEN (cuadro de mando)
# ==========================================================================
ws = ws_res
title(ws, "Modelo financiero — Subasta SUB-JA-2026-265709",
      "Plaza de aparcamiento para moto nº 10 · C/ Virtud 38 – C/ Constancia 71 · López de Gomara (Triana), Sevilla")
put(ws, "B4", "Escenario activo:", F_LBL_B)
put(ws, "C4", "=NombreEscenario", F_LINK_B, fill=FILL_KEY, border=BOX)
put(ws, "E4", "Estrategia principal:", F_LBL_B)
put(ws, "F4", "=NombreEstrategia", F_LINK_B, fill=FILL_KEY, border=BOX)
put(ws, "H4", "Cambiar en Supuestos (F7 y F9)", F_NOTE)


def panel(col_l, col_v, r0, ttl, items):
    section(ws, r0, ttl, col_l, col_v)
    rr = r0 + 1
    for lbl, fx, fmt, bold in items:
        put(ws, f"{col_l}{rr}", lbl, F_LBL_B if bold else F_LBL)
        put(ws, f"{col_v}{rr}", fx, F_LINK_B if bold else F_LINK, fmt=fmt, border=BOX,
            fill=FILL_KPI if bold else None, align=Alignment(horizontal="right"))
        rr += 1
    return rr


r1 = panel("B", "C", 6, "La subasta", [
    ("Valor de subasta (= tasación)", "=ValorSubasta", EUR2, True),
    ("Depósito para pujar (5%)", "=Deposito", EUR2, False),
    ("70% del valor (aprobación directa)", "=UmbralSetenta", EUR2, False),
    ("50% del valor", "=UmbralCincuenta", EUR2, False),
    ("Cantidad reclamada por la Comunidad", "=CantidadReclamada", EUR2, False),
    ("Cargas anteriores que asume el comprador", "=CargasAnteriores", EUR2, False),
    ("Superficie construida / útil (m²)", '=SupConstruida&" / "&SupUtil', None, False),
    ("Fecha de cierre (estimada, verificar)", "=FechaCierre", DATE, False),
])
r2 = panel("E", "F", 6, "Precio de compra (escenario activo)", [
    ("Puja (precio de remate)", "=Puja", EUR0, True),
    ("Puja / valor de subasta", "=Puja/ValorSubasta", PCT1, False),
    ("Impuesto de adquisición", "=ImpuestoAdq", EUR0, False),
    ("Costes fijos (registro, gestoría, deudas, posesión)", "=CostesFijosAdq", EUR0, False),
    ("Inversión total — compraventa", "=FlipInversion", EUR0, True),
    ("Inversión total — alquiler", "=AlqInversion", EUR0, False),
    ("Sobrecoste si tributa por IVA en vez de ITP", "=ImpactoIVA", EUR0, False),
    ("Régimen art. 670 LEC de la puja",
     '=IF(Puja>=UmbralSetenta,"Aprobación directa",IF(Puja>=UmbralCincuenta,"Mejorable (10 + 5 días)","Decide el LAJ"))', None, False),
])
r3 = panel("B", "C", max(r1, r2) + 1, "Estrategia 1 — Compraventa", [
    ("Precio de cierre esperado", "=PrecioCierre", EUR0, False),
    ("Ingreso neto por la venta", "=FlipVentaNeta", EUR0, False),
    ("Beneficio antes de impuestos", "=FlipBAI", EUR0, True),
    ("Beneficio neto", "=FlipNeto", EUR0, False),
    ("ROI antes de impuestos", "=FlipROI", PCT1, True),
    ("TIR anual antes de impuestos", "=FlipTIR", PCT1, False),
    ("Plazo (meses)", "=MesVenta", IDX, False),
])
r4 = panel("E", "F", max(r1, r2) + 1, "Estrategia 2 — Alquiler", [
    ("Renta de mercado (€/mes)", "=Renta", EUR0, False),
    ("Renta neta de IVA (€/mes)", "=RentaNeta", EUR2, False),
    ("NOI año 1", "=AlqNOI", EUR0, False),
    ("Rentabilidad bruta", "=AlqRentBruta", PCT1, False),
    ("Rentabilidad neta", "=AlqRentNeta", PCT1, True),
    ("TIR a 10 años (con venta)", "=AlqTIR", PCT1, True),
    ("VAN a la tasa de descuento", "=AlqVAN", EUR0, False),
])
r5 = panel("B", "C", max(r3, r4) + 1, "Puja máxima", [
    ("Puja mínima prudente (70% del valor)", "=PujaMinPrudente", EUR0, False),
    ("Puja máxima — compraventa (ROI objetivo)", "=PujaMaxFlip", EUR0, True),
    ("Puja máxima — alquiler (rentabilidad objetivo)", "=PujaMaxAlq", EUR0, False),
    ("Puja de equilibrio — compraventa (beneficio 0)", "=PujaEquilibrio", EUR0, False),
    ("Puja máxima recomendada (estrategia principal)", "=PujaMaxSel", EUR0, True),
])
r6 = panel("E", "F", max(r3, r4) + 1, "Escenarios (beneficio y ROI de la compraventa)", [
    ("Pesimista — beneficio / ROI", "=EscBAI_P", EUR0, False),
    ("", "=EscROI_P", PCT1, False),
    ("Base — beneficio / ROI", "=EscBAI_B", EUR0, True),
    ("", "=EscROI_B", PCT1, True),
    ("Optimista — beneficio / ROI", "=EscBAI_O", EUR0, False),
    ("", "=EscROI_O", PCT1, False),
])
rv = max(r5, r6) + 1
section(ws, rv, "Veredicto sobre la puja del escenario activo", "B", "F")
put(ws, f"B{rv+1}",
    '=IF(Puja<UmbralSetenta,"ATENCIÓN: la puja es inferior al 70% del valor de subasta. La Comunidad ejecutante podría adjudicarse la plaza (art. 670.4 LEC).",'
    'IF(Puja>PujaMaxSel,"NO RECOMENDADA: la puja supera la puja máxima y no alcanza la rentabilidad objetivo de la estrategia principal.",'
    '"RECOMENDADA: la puja está entre el 70% del valor de subasta y la puja máxima de la estrategia principal."))',
    Font(name=FONT, size=11, bold=True, color=BLACK), align=WRAP)
ws.merge_cells(f"B{rv+1}:F{rv+1}")
ws.row_dimensions[rv + 1].height = 34
ws.conditional_formatting.add(f"B{rv+1}", FormulaRule(formula=[f'LEFT($B${rv+1},4)="RECO"'], fill=FILL_OK))
ws.conditional_formatting.add(f"B{rv+1}", FormulaRule(formula=[f'LEFT($B${rv+1},4)="ATEN"'], fill=FILL_WARN))
ws.conditional_formatting.add(f"B{rv+1}", FormulaRule(formula=[f'LEFT($B${rv+1},2)="NO"'], fill=FILL_BAD))
put(ws, f"B{rv+2}", "Rango de puja recomendado:", F_LBL_B)
put(ws, f"C{rv+2}", "=PujaMinPrudente", F_LINK_B, fmt=EUR0, border=BOX)
put(ws, f"D{rv+2}", "a", F_LBL, align=CENTER)
put(ws, f"E{rv+2}", "=MAX(PujaMinPrudente,PujaMaxSel)", F_FX_B, fmt=EUR0, border=BOX, align=Alignment(horizontal="right"))
put(ws, f"F{rv+2}", '=IF(PujaMaxSel<PujaMinPrudente,"La puja máxima no llega al 70%: no pujar","")', F_NOTE)
put(ws, f"B{rv+3}", "Controles del modelo:", F_LBL_B)
put(ws, f"C{rv+3}", '=IF(AND(FlipControl="OK",AlqControl="OK"),"OK","REVISAR")', F_FX_B, border=BOX, align=CENTER)

rk = rv + 5
section(ws, rk, "Alertas clave", "B", "F")
alerts = [
    "1. La plaza solo admite una moto (3,75 m² útiles): limita la demanda de reventa y de alquiler.",
    "2. Hay que pujar al menos el 70% del valor (4.954,92 €); por debajo, la Comunidad ejecutante, a la que se deben 10.083,73 €, puede adjudicársela.",
    "3. Riesgo fiscal: si la adjudicación es una 'primera entrega' del promotor, tributa por IVA 21% en lugar de ITP 7%.",
    "4. Las cuotas de comunidad impagadas (año en curso y 3 anteriores) y el IBI pendiente pasan al comprador: pedir certificados antes de pujar.",
    "5. Pendiente de la ficha del Portal: fecha exacta de cierre, tramos entre pujas, pujas actuales y subastas relacionadas.",
]
for i, a in enumerate(alerts):
    put(ws, f"B{rk+1+i}", a, F_LBL, align=WRAP)
    ws.merge_cells(f"B{rk+1+i}:F{rk+1+i}")
    ws.row_dimensions[rk + 1 + i].height = 28
put(ws, f"B{rk+7}", "Cifras en euros. Celdas verdes = vínculos a otras hojas; los supuestos se editan en la hoja Supuestos.", F_NOTE)
widths(ws, {"A": 2, "B": 48, "C": 17, "D": 4, "E": 46, "F": 22, "G": 2, "H": 30})
setup_print(ws, landscape=False)

# Propiedades del libro
wb.properties.title = "Modelo financiero SUB-JA-2026-265709"
wb.properties.subject = "Subasta judicial — plaza de aparcamiento (moto), C/ Virtud 38 – C/ Constancia 71, Sevilla"
wb.properties.creator = "Sem2a"
for s in wb.worksheets:
    s.sheet_properties.tabColor = {"Resumen": NAVY, "Supuestos": "FFC000", "Subasta": "70AD47"}.get(s.title, "A6A6A6")
wb.active = 0
wb.save(OUT)
print("OK ->", OUT)
